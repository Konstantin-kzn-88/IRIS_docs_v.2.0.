from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.report_component_fatality_risk import (
    NOT_CALCULATED,
    ReportComponentFatalityRiskError,
    load_component_fatality_risk_rows,
)


MARKER = "{{NGK_BACKGROUND_RISK_COMPARISON}}"
BACKGROUND_VALUES = (
    ("Нефтегазодобывающая промышленность", -4.3, 7.3),
    ("Нефтеперерабатывающая промышленность", -5.8, 5.2),
    ("Нефтехимическая промышленность", -8.9, 2.4),
    ("Объекты газораспределения и газопотребления*", -8.7, 2.6),
    ("Магистральный трубопроводный транспорт*", -10.9, 1.6),
)
SOURCE_NOTE = (
    "Примечание — * только для работников. Фоновые значения за 2013–2022 гг. "
    "приняты по таблице № 1 приложения № 2 к приказу Ростехнадзора "
    "от 12.09.2023 № 331."
)


class ReportNgkBackgroundRiskError(Exception):
    pass


def _decimal(value: float, digits: int = 1) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def _comparison(calculated: float, background: float) -> str:
    if calculated == 0:
        return "Ниже фонового значения"
    if calculated == background:
        return "Соответствует фоновому значению"
    if calculated < background:
        return f"Ниже в {_decimal(background / calculated)} раза"
    return f"Выше в {_decimal(calculated / background)} раза"


def load_ngk_background_risk_rows(
    project_directory: Path | str,
) -> tuple[str, tuple[dict[str, str], ...]]:
    try:
        component_rows = load_component_fatality_risk_rows(project_directory)
    except ReportComponentFatalityRiskError as exc:
        raise ReportNgkBackgroundRiskError(str(exc)) from exc

    individual_values = [
        float(item["individual"])
        for item in component_rows
        if item["individual"] != NOT_CALCULATED
    ]
    maximum = max(individual_values) if individual_values else None
    maximum_text = NOT_CALCULATED if maximum is None else f"{maximum:.3E}"
    rows = tuple(
        {
            "industry": industry,
            "dbr": _decimal(dbr),
            "per_100k": _decimal(per_100k),
            "per_year": f"{per_100k / 100_000:.3E}",
            "comparison": (
                NOT_CALCULATED
                if maximum is None
                else _comparison(maximum, per_100k / 100_000)
            ),
        }
        for industry, dbr, per_100k in BACKGROUND_VALUES
    )
    return maximum_text, rows


def _set_font(run: Any, size: float) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_text(
    cell: Any,
    value: str,
    *,
    bold: bool = False,
    centered: bool = False,
    font_size: float = 8.5,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    _set_font(run, font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_paragraph_text(
    paragraph: Any,
    value: str,
    *,
    bold: bool = False,
    font_size: float = 9,
    space_before: float = 0,
    space_after: float = 0,
) -> None:
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(value)
    run.bold = bold
    _set_font(run, font_size)


def _paragraph_section(document: DocumentType, paragraph_element: Any) -> Any:
    section_index = 0
    for child in document.element.body:
        if child is paragraph_element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            section_index += 1
    return document.sections[min(section_index, len(document.sections) - 1)]


def _set_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.30, 0.13, 0.18, 0.15, 0.24)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def render_ngk_background_risk_comparison(
    document: DocumentType,
    maximum_risk: str,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False
    section = _paragraph_section(document, marker_paragraph._p)
    lead = document.add_paragraph()
    marker_paragraph._p.addnext(lead._p)
    _set_paragraph_text(
        lead,
        "Максимальный расчётный индивидуальный риск гибели по составляющим "
        f"ОПО: {maximum_risk} 1/год.",
        bold=True,
        space_after=3,
    )
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    lead._p.addnext(table._tbl)
    note = document.add_paragraph()
    table._tbl.addnext(note._p)
    _set_paragraph_text(note, SOURCE_NOTE, font_size=7.5, space_before=2)
    headers = (
        "Отрасль нефтегазового комплекса",
        "Уровень риска, дБR",
        "Погибших на 100 тыс. рискующих",
        "Фоновый риск, 1/год",
        "Сравнение с расчётным риском",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=8)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for item in rows:
        cells = table.add_row().cells
        values = (
            item["industry"],
            item["dbr"],
            item["per_100k"],
            item["per_year"],
            item["comparison"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column in (1, 2, 3))
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_table_geometry(section, table)
    return True


__all__ = [
    "BACKGROUND_VALUES",
    "MARKER",
    "ReportNgkBackgroundRiskError",
    "load_ngk_background_risk_rows",
    "render_ngk_background_risk_comparison",
]
