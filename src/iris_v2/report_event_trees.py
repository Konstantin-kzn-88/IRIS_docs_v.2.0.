import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

from iris_v2.calculation_cases import FILE_NAME as CASES_FILE_NAME
from iris_v2.event_tree_images import EventTreeImageError, EventTreeImageService
from iris_v2.typical_scenarios import TypicalScenarioError, TypicalScenarioService


MARKER = "{{EVENT_TREES_SECTION}}"


class ReportEventTreesError(Exception):
    pass


@dataclass(frozen=True)
class EventTreeReportItem:
    equipment_type: int
    kind: int
    equipment_name: str
    kind_name: str
    path: Path


def _load_cases(project: Path) -> list[dict[str, Any]]:
    path = project / CASES_FILE_NAME
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ReportEventTreesError(
            "Расчётные сценарии не сформированы. "
            "Сначала выполните модуль «Расчётные сценарии»"
        ) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportEventTreesError(
            f"Не удалось прочитать {CASES_FILE_NAME}"
        ) from exc
    cases = raw.get("cases") if isinstance(raw, dict) else None
    if (
        not isinstance(cases, list)
        or not cases
        or any(not isinstance(item, dict) for item in cases)
    ):
        raise ReportEventTreesError(
            f"Файл {CASES_FILE_NAME} не содержит расчётных сценариев"
        )
    return cases


def prepare_event_trees(
    project_directory: Path | str,
) -> tuple[EventTreeReportItem, ...]:
    project = Path(project_directory)
    cases = _load_cases(project)
    pairs: list[tuple[int, int]] = []
    seen: set[tuple[int, int]] = set()
    for index, case in enumerate(cases, start=1):
        equipment_type = case.get("equipment_type")
        kind = case.get("kind")
        if (
            isinstance(equipment_type, bool)
            or not isinstance(equipment_type, int)
            or equipment_type not in range(10)
        ):
            raise ReportEventTreesError(
                f"Сценарий {index}: equipment_type должен быть от 0 до 9"
            )
        if isinstance(kind, bool) or not isinstance(kind, int) or kind not in range(10):
            raise ReportEventTreesError(
                f"Сценарий {index}: kind должен быть от 0 до 9"
            )
        pair = (equipment_type, kind)
        if pair not in seen:
            seen.add(pair)
            pairs.append(pair)

    try:
        catalog = TypicalScenarioService().load()
        copied = EventTreeImageService().copy_to_project(
            project, pairs=set(pairs)
        )
    except (TypicalScenarioError, EventTreeImageError) as exc:
        raise ReportEventTreesError(str(exc)) from exc
    paths = {path.name: path for path in copied.files}
    result: list[EventTreeReportItem] = []
    for equipment_type, kind in pairs:
        filename = f"event_tree_eq{equipment_type:02d}_kind{kind:02d}.png"
        path = paths.get(filename)
        if path is None:
            raise ReportEventTreesError(f"Не найдено дерево событий: {filename}")
        result.append(
            EventTreeReportItem(
                equipment_type=equipment_type,
                kind=kind,
                equipment_name=catalog.equipment_types[equipment_type],
                kind_name=catalog.kinds[kind],
                path=path,
            )
        )
    return tuple(result)


def _paragraph_section(document: DocumentType, element: Any) -> Any:
    section_index = 0
    for child in document.element.body:
        if child is element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            section_index += 1
    return document.sections[min(section_index, len(document.sections) - 1)]


def _set_font(run: Any, size: float, *, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _short_kind_name(value: str) -> str:
    base, separator, suffix = value.rpartition(" (")
    if not separator or not suffix.endswith(")"):
        return value
    parenthetical = suffix[:-1].strip()
    if parenthetical.isupper() and len(parenthetical) <= 10:
        return parenthetical
    return base.strip()


def _picture_width(path: Path, available_width: int, max_height: int) -> int:
    try:
        with Image.open(path) as image:
            width_px, height_px = image.size
    except (OSError, ValueError) as exc:
        raise ReportEventTreesError(
            f"Не удалось прочитать PNG-дерево: {path.name}"
        ) from exc
    if width_px <= 0 or height_px <= 0:
        raise ReportEventTreesError(f"Некорректный PNG-файл: {path.name}")
    return min(available_width, int(max_height * width_px / height_px))


def _append_field(
    paragraph: Any,
    instruction: str,
    display_text: str,
    *,
    bookmark_name: str | None = None,
    bookmark_id: int | None = None,
) -> None:
    if bookmark_name is not None and bookmark_id is not None:
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(bookmark_id))
        bookmark_start.set(qn("w:name"), bookmark_name)
        paragraph._p.append(bookmark_start)

    begin_run = paragraph.add_run()
    _set_font(begin_run, 11)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    _set_font(instruction_run, 11)
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = f" {instruction} "
    instruction_run._r.append(instruction_element)

    separate_run = paragraph.add_run()
    _set_font(separate_run, 11)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    result_run = paragraph.add_run(display_text)
    _set_font(result_run, 11)

    end_run = paragraph.add_run()
    _set_font(end_run, 11)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)

    if bookmark_name is not None and bookmark_id is not None:
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.append(bookmark_end)


def _fields_in_paragraph(paragraph: Any) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for element in paragraph.iter():
        if element.tag == qn("w:fldChar"):
            field_type = element.get(qn("w:fldCharType"))
            if field_type == "begin":
                current = {"instruction": "", "results": [], "separated": False}
            elif current is not None and field_type == "separate":
                current["separated"] = True
            elif current is not None and field_type == "end":
                fields.append(current)
                current = None
        elif current is not None and element.tag == qn("w:instrText"):
            current["instruction"] += element.text or ""
        elif (
            current is not None
            and current["separated"]
            and element.tag == qn("w:t")
        ):
            current["results"].append(element)
    return fields


def _set_field_result(field: dict[str, Any], value: int) -> None:
    results = field["results"]
    if not results:
        return
    results[0].text = str(value)
    for extra in results[1:]:
        extra.text = ""


def _materialize_figure_fields(document: DocumentType) -> None:
    paragraphs: list[tuple[Any, list[dict[str, Any]]]] = []
    bookmarks: dict[str, int] = {}
    figure_number = 0
    for paragraph in document.element.body.iter(qn("w:p")):
        fields = _fields_in_paragraph(paragraph)
        paragraphs.append((paragraph, fields))
        for field in fields:
            if not re.search(
                r"\bSEQ\s+Рисунок\b", field["instruction"], re.IGNORECASE
            ):
                continue
            figure_number += 1
            _set_field_result(field, figure_number)
            for bookmark in paragraph.iter(qn("w:bookmarkStart")):
                name = bookmark.get(qn("w:name"))
                if name:
                    bookmarks[name] = figure_number

    for _paragraph, fields in paragraphs:
        for field in fields:
            match = re.search(
                r"\bREF\s+([^\s\\]+)", field["instruction"], re.IGNORECASE
            )
            if match and match.group(1) in bookmarks:
                _set_field_result(field, bookmarks[match.group(1)])

    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _next_bookmark_id(document: DocumentType) -> int:
    values = []
    for bookmark in document.element.body.iter(qn("w:bookmarkStart")):
        value = bookmark.get(qn("w:id"), "")
        if value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def render_event_trees_section(
    document: DocumentType,
    items: tuple[EventTreeReportItem, ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False
    if not items:
        raise ReportEventTreesError("Не выбрано ни одного дерева событий")
    section = _paragraph_section(document, marker_paragraph._p)
    available_width = section.page_width - section.left_margin - section.right_margin
    max_height = Inches(5.5)
    anchor = marker_paragraph._p
    bookmark_id = _next_bookmark_id(document)
    bookmark_names = tuple(
        f"IrisEventTreeFigure{index}" for index in range(1, len(items) + 1)
    )
    lead = document.add_paragraph()
    lead.paragraph_format.keep_with_next = True
    lead.paragraph_format.space_before = Pt(0)
    lead.paragraph_format.space_after = Pt(0)
    lead.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if len(items) == 1:
        _set_font(lead.add_run("Дерево событий представлено на рисунке "), 11)
        _append_field(lead, f"REF {bookmark_names[0]} \\h", "0")
    else:
        _set_font(
            lead.add_run("Деревья событий представлены на рисунках "), 11
        )
        _append_field(lead, f"REF {bookmark_names[0]} \\h", "0")
        _set_font(lead.add_run("–"), 11)
        _append_field(lead, f"REF {bookmark_names[-1]} \\h", "0")
    _set_font(lead.add_run("."), 11)
    anchor.addnext(lead._p)
    anchor = lead._p

    for index, item in enumerate(items, start=1):
        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.page_break_before = True
        picture_paragraph.paragraph_format.keep_with_next = True
        picture_paragraph.paragraph_format.space_before = Pt(0)
        picture_paragraph.paragraph_format.space_after = Pt(3)
        picture_run = picture_paragraph.add_run()
        picture_run.add_picture(
            str(item.path),
            width=_picture_width(item.path, available_width, max_height),
        )
        anchor.addnext(picture_paragraph._p)
        anchor = picture_paragraph._p

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_together = True
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(6)
        _set_font(caption.add_run("Рисунок "), 11)
        _append_field(
            caption,
            "SEQ Рисунок \\* ARABIC",
            "0",
            bookmark_name=bookmark_names[index - 1],
            bookmark_id=bookmark_id + index - 1,
        )
        run = caption.add_run(
            " – Дерево событий для типа оборудования "
            f"«{item.equipment_name}» и вида опасного вещества "
            f"«{_short_kind_name(item.kind_name)}»"
        )
        _set_font(run, 11)
        anchor.addnext(caption._p)
        anchor = caption._p
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _materialize_figure_fields(document)
    return True
