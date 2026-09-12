from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.substances import KIND_NAMES


MARKER = "{{SUBSTANCES_INFO_SECTION}}"


class ReportSubstancesInfoError(Exception):
    pass


def _is_value(value: Any) -> bool:
    return value is not None and not isinstance(value, bool) and value != ""


def _number(value: Any) -> str:
    return format(value, ".12g").replace(".", ",")


def _property(value: Any, label: str, unit: str) -> str | None:
    if not _is_value(value):
        return None
    text = _number(value) if isinstance(value, (int, float)) else str(value).strip()
    return f"{label} {text} {unit}".strip()


def short_substance_characteristic(item: dict[str, Any]) -> str:
    kind = item.get("kind")
    parts = [KIND_NAMES[kind]] if kind in KIND_NAMES else []
    physical = item.get("physical") if isinstance(item.get("physical"), dict) else {}
    explosion = (
        item.get("explosion") if isinstance(item.get("explosion"), dict) else {}
    )
    toxicity = (
        item.get("toxicity") if isinstance(item.get("toxicity"), dict) else {}
    )
    properties = [
        _property(
            physical.get("density_liquid_kg_per_m3"),
            "плотность жидкости",
            "кг/м³",
        ),
        _property(physical.get("boiling_point_C"), "температура кипения", "°C"),
        _property(explosion.get("flash_point_C"), "температура вспышки", "°C"),
        _property(
            explosion.get("autoignition_temp_C"),
            "температура самовоспламенения",
            "°C",
        ),
        _property(toxicity.get("hazard_class"), "класс опасности", ""),
        _property(toxicity.get("pdk_mg_per_m3"), "ПДК", "мг/м³"),
    ]
    properties = [value for value in properties if value]
    if properties:
        parts.append("Основные свойства: " + "; ".join(properties))
    notes = str(item.get("notes", "")).strip()
    if notes and notes != "-":
        parts.append(notes.rstrip("."))
    impact = str(item.get("impact", "")).strip()
    if impact and impact != "-":
        parts.append("Воздействие на людей: " + impact.rstrip("."))
    neutralization = str(item.get("neutralization_methods", "")).strip()
    if neutralization and neutralization != "-":
        parts.append("Способы обезвреживания: " + neutralization.rstrip("."))
    first_aid = str(item.get("first_aid", "")).strip()
    if first_aid and first_aid != "-":
        parts.append("Первая помощь: " + first_aid.rstrip("."))
    return ". ".join(parts).strip() + ("." if parts else "Сведения не заполнены.")


def _set_repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def _prevent_row_split(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _set_cell_text(
    cell: Any,
    value: str,
    *,
    bold: bool = False,
    centered: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_table_geometry(document: DocumentType, table: Any) -> None:
    section = document.sections[0]
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    first_width = int(total_twips * 0.28)
    widths = (first_width, total_twips - first_width)
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def render_substances_info_section(
    document: DocumentType,
    substances: tuple[dict[str, Any], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False
    if not substances:
        raise ReportSubstancesInfoError(
            "Вещества проекта не выбраны. Сначала заполните модуль «Вещества»"
        )

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    for cell, value in zip(
        table.rows[0].cells,
        ("Наименование вещества", "Краткая характеристика"),
    ):
        _set_cell_text(cell, value, bold=True, centered=True)
    _set_repeat_header(table.rows[0])

    for item in substances:
        cells = table.add_row().cells
        _set_cell_text(cells[0], str(item.get("name", "—")).strip() or "—")
        _set_cell_text(cells[1], short_substance_characteristic(item))

    for row in table.rows:
        _prevent_row_split(row)
    _set_table_geometry(document, table)
    return True


__all__ = [
    "MARKER",
    "ReportSubstancesInfoError",
    "render_substances_info_section",
    "short_substance_characteristic",
]
