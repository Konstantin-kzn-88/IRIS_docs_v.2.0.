import json
import math
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.risk_calculation import FILE_NAME as RISK_FILE_NAME
from iris_v2.risk_summary import FILE_NAME as SUMMARY_FILE_NAME


MARKER = "{{FATALITY_RISK_BY_COMPONENT_SECTION}}"
NOT_CALCULATED = "Не рассчитан"


class ReportComponentFatalityRiskError(Exception):
    pass


def _load_object(path: Path, missing_message: str) -> dict[str, Any]:
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ReportComponentFatalityRiskError(missing_message) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportComponentFatalityRiskError(
            f"Не удалось прочитать {path.name}"
        ) from exc
    if not isinstance(raw, dict):
        raise ReportComponentFatalityRiskError(
            f"Файл {path.name} повреждён: ожидается объект"
        )
    return raw


def _number(value: Any, label: str) -> float:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or float(value) < 0
    ):
        raise ReportComponentFatalityRiskError(
            f"{label} должно быть числом не меньше нуля"
        )
    return float(value)


def _count(value: Any, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ReportComponentFatalityRiskError(
            f"{label} должно быть целым числом не меньше нуля"
        )
    return value


def _same(left: float, right: float) -> bool:
    return math.isclose(left, right, rel_tol=1e-12, abs_tol=0.0)


def load_component_fatality_risk_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    project = Path(project_directory)
    risk = _load_object(
        project / RISK_FILE_NAME,
        "Риски не рассчитаны. Сначала выполните модуль «Расчёт риска»",
    )
    summary = _load_object(
        project / SUMMARY_FILE_NAME,
        "Сводные показатели риска не рассчитаны. "
        "Сначала выполните модуль «Сводные показатели риска»",
    )
    results = risk.get("results")
    if not isinstance(results, list) or not results:
        raise ReportComponentFatalityRiskError(
            f"Файл {RISK_FILE_NAME} повреждён: results должен быть непустым списком"
        )

    grouped: dict[str, dict[str, Any]] = {}
    codes: set[str] = set()
    for index, item in enumerate(results, start=1):
        if not isinstance(item, dict):
            raise ReportComponentFatalityRiskError(
                f"{RISK_FILE_NAME}, запись {index}: ожидается объект"
            )
        code = str(item.get("scenario_code", "")).strip()
        component = str(item.get("hazard_component", "")).strip()
        if not code or code in codes:
            raise ReportComponentFatalityRiskError(
                f"{RISK_FILE_NAME}, запись {index}: "
                "пустой или повторяющийся scenario_code"
            )
        if not component:
            raise ReportComponentFatalityRiskError(
                f"Сценарий {code}: не заполнена составляющая ОПО"
            )
        codes.add(code)
        fatalities = _count(
            item.get("fatalities_count"), f"Сценарий {code}: fatalities_count"
        )
        frequency = _number(
            item.get("scenario_frequency"), f"Сценарий {code}: scenario_frequency"
        )
        collective = _number(
            item.get("collective_risk_fatalities"),
            f"Сценарий {code}: collective_risk_fatalities",
        )
        individual_value = item.get("individual_risk_fatalities")
        individual = (
            None
            if individual_value is None
            else _number(
                individual_value,
                f"Сценарий {code}: individual_risk_fatalities",
            )
        )
        values = grouped.setdefault(
            component,
            {
                "count": 0,
                "max_fatalities": 0,
                "collective": 0.0,
                "individual": 0.0,
                "has_individual": False,
                "fatal_frequency": 0.0,
            },
        )
        values["count"] += 1
        values["max_fatalities"] = max(values["max_fatalities"], fatalities)
        values["collective"] += collective
        if individual is not None:
            values["individual"] += individual
            values["has_individual"] = True
        if fatalities >= 1:
            values["fatal_frequency"] += frequency

    if summary.get("case_count") != len(results):
        raise ReportComponentFatalityRiskError(
            f"Файл {SUMMARY_FILE_NAME} устарел: не совпадает case_count"
        )
    if summary.get("component_count") != len(grouped):
        raise ReportComponentFatalityRiskError(
            f"Файл {SUMMARY_FILE_NAME} устарел: не совпадает component_count"
        )
    if summary.get("risk_unit") != "1/год":
        raise ReportComponentFatalityRiskError(
            f"Файл {SUMMARY_FILE_NAME} содержит неизвестную единицу риска"
        )
    components = summary.get("components")
    if not isinstance(components, list) or len(components) != len(grouped):
        raise ReportComponentFatalityRiskError(
            f"Файл {SUMMARY_FILE_NAME} устарел: неверный перечень составляющих ОПО"
        )

    rows: list[dict[str, str]] = []
    for index, ((name, expected), stored) in enumerate(
        zip(grouped.items(), components), start=1
    ):
        if not isinstance(stored, dict) or stored.get("hazard_component") != name:
            raise ReportComponentFatalityRiskError(
                f"Файл {SUMMARY_FILE_NAME}, составляющая {index}: неверное наименование"
            )
        if stored.get("scenario_count") != expected["count"]:
            raise ReportComponentFatalityRiskError(
                f"Файл {SUMMARY_FILE_NAME}, составляющая «{name}»: "
                "не совпадает число сценариев"
            )
        collective = _number(
            stored.get("collective_risk_fatalities"),
            f"Составляющая «{name}»: collective_risk_fatalities",
        )
        fatal_frequency = _number(
            stored.get("fatal_accident_frequency"),
            f"Составляющая «{name}»: fatal_accident_frequency",
        )
        individual_value = stored.get("individual_risk_fatalities")
        individual = (
            None
            if individual_value is None
            else _number(
                individual_value,
                f"Составляющая «{name}»: individual_risk_fatalities",
            )
        )
        expected_individual = (
            expected["individual"] if expected["has_individual"] else None
        )
        if (
            not _same(collective, expected["collective"])
            or not _same(fatal_frequency, expected["fatal_frequency"])
            or (individual is None) != (expected_individual is None)
            or (
                individual is not None
                and expected_individual is not None
                and not _same(individual, expected_individual)
            )
        ):
            raise ReportComponentFatalityRiskError(
                f"Файл {SUMMARY_FILE_NAME} устарел: риск гибели по составляющей "
                f"«{name}» не совпадает с {RISK_FILE_NAME}"
            )
        rows.append(
            {
                "component": name,
                "max_fatalities": str(expected["max_fatalities"]),
                "collective": f"{collective:.3E}",
                "individual": (
                    NOT_CALCULATED if individual is None else f"{individual:.3E}"
                ),
                "fatal_frequency": f"{fatal_frequency:.3E}",
            }
        )
    return tuple(rows)


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_text(
    cell: Any,
    value: str,
    *,
    bold: bool = False,
    centered: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _paragraph_section(document: DocumentType, paragraph_element: Any) -> Any:
    section_index = 0
    for child in document.element.body:
        if child is paragraph_element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            section_index += 1
    return document.sections[min(section_index, len(document.sections) - 1)]


def _set_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.29, 0.15, 0.19, 0.19, 0.18)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def render_component_fatality_risk_section(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False
    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Максимальное число погибших, чел.",
        "Коллективный риск гибели, чел./год",
        "Индивидуальный риск гибели, 1/год",
        "Частота аварий с гибелью, 1/год",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for item in rows:
        cells = table.add_row().cells
        values = (
            item["component"],
            item["max_fatalities"],
            item["collective"],
            item["individual"],
            item["fatal_frequency"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column > 0)
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_table_geometry(section, table)
    return True


__all__ = [
    "MARKER",
    "NOT_CALCULATED",
    "ReportComponentFatalityRiskError",
    "load_component_fatality_risk_rows",
    "render_component_fatality_risk_section",
]
