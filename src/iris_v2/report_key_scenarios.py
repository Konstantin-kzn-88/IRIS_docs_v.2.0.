import json
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.key_scenarios import KeyScenariosError, KeyScenariosService
from iris_v2.report_damage import DAMAGE_FIELDS, ReportDamageError, load_damage_rows
from iris_v2.report_impact_zones import (
    ZONE_FIELDS,
    ReportImpactZonesError,
    load_impact_zone_rows,
)
from iris_v2.report_ov_amount import ReportOvAmountError, load_ov_amount_rows
from iris_v2.impact_zones import FILE_NAME as IMPACT_ZONES_FILE_NAME


MARKER = "{{TOP_SCENARIOS_BY_COMPONENT_SECTION}}"
DESCRIPTION_MARKER = "{{TOP_SCENARIOS_DESC_BY_COMPONENT}}"
PF_MARKER = "{{TOP_SCENARIOS_PF_BY_COMPONENT}}"
PEOPLE_MARKER = "{{TOP_SCENARIOS_FATALITIES_INJURED}}"
DAMAGE_MARKER = "{{TOP_SCENARIOS_DAMAGE}}"
CONCLUSION_MARKER = "{{TOP_SCENARIOS_FINAL_CONCLUSION}}"
ACCIDENT_DESCRIPTION_MARKER = "{{SITUATION_PLAN_ACCIDENTS_TABLE}}"

CALCULATION_METHODS = {
    0: "—",
    1: "Приказ МЧС России от 10.07.2009 № 404",
    2: "СП 12.13130.2009",
    3: "Приказ МЧС России от 10.07.2009 № 404",
    4: "Временная оценка по массе",
    5: "Приказ МЧС России от 10.07.2009 № 404, формулы П3.71–П3.72",
    6: "Приказ МЧС России от 10.07.2009 № 404",
    7: "Расчёт площади химически опасного пролива",
}

ZONE_DESCRIPTIONS = {
    "q_10_5_m": "зона теплового излучения с интенсивностью 10,5 кВт/м²",
    "q_7_0_m": "зона теплового излучения с интенсивностью 7,0 кВт/м²",
    "q_4_2_m": "зона теплового излучения с интенсивностью 4,2 кВт/м²",
    "q_1_4_m": "зона теплового излучения с интенсивностью 1,4 кВт/м²",
    "p_100_m": "зона полного разрушения зданий (100 кПа)",
    "p_70_m": "зона сильных разрушений зданий (70 кПа)",
    "p_28_m": "зона средних разрушений зданий (28 кПа)",
    "p_14_m": "зона умеренных разрушений зданий (14 кПа)",
    "p_5_m": "зона слабых разрушений зданий (5 кПа)",
    "p_2_m": "зона разрушения остекления (2 кПа)",
    "jet_fire_length_m": "длина факела",
    "jet_fire_diameter_m": "диаметр факела",
    "lel_radius_m": "радиус нижнего концентрационного предела распространения пламени",
    "flash_fire_radius_m": "радиус пожара-вспышки",
    "lethal_radius_m": "радиус смертельной токсодозы (LD)",
    "threshold_radius_m": "радиус пороговой токсодозы (PD)",
    "dose_600_m": "зона тепловой дозы 600 кДж/м²",
    "dose_320_m": "зона тепловой дозы 320 кДж/м²",
    "dose_220_m": "зона тепловой дозы 220 кДж/м²",
    "dose_120_m": "зона тепловой дозы 120 кДж/м²",
    "spill_area_m2": "площадь химически опасного пролива",
}


class ReportKeyScenariosError(Exception):
    pass


def load_key_scenario_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    try:
        result = KeyScenariosService().calculate(project_directory)
    except KeyScenariosError as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    return tuple(
        {
            "component": str(item["hazard_component"]),
            "scenario_type": str(item["scenario_type_name"]),
            "scenario_code": str(item["scenario_code"]),
            "equipment": str(item["equipment_name"]),
            "fatalities": str(item["fatalities_count"]),
            "injured": str(item["injured_count"]),
            "affected": str(item["fatalities_count"] + item["injured_count"]),
            "damage": f"{float(item['total_damage']):.1f}".replace(".", ","),
            "frequency": f"{float(item['scenario_frequency']):.3E}",
        }
        for item in result.rows
    )


def load_key_scenario_description_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    try:
        result = KeyScenariosService().calculate(project_directory)
    except KeyScenariosError as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    return tuple(
        {
            "component": str(item["hazard_component"]),
            "scenario_type": str(item["scenario_type_name"]),
            "scenario_code": str(item["scenario_code"]),
            "description": str(item["scenario_text"]),
        }
        for item in result.rows
    )


def load_key_scenario_pf_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    try:
        selected = KeyScenariosService().calculate(project_directory).rows
        impact_rows = load_impact_zone_rows(project_directory)
    except (KeyScenariosError, ReportImpactZonesError) as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    impact_by_code = {item["code"]: item for item in impact_rows}
    rows: list[dict[str, str]] = []
    for item in selected:
        code = str(item["scenario_code"])
        impact = impact_by_code.get(code)
        if impact is None:
            raise ReportKeyScenariosError(
                f"В impact_zones.json отсутствует ключевой сценарий {code}"
            )
        expected_equipment = (
            f"{item['equipment_name']} ({item['hazard_component']})"
        )
        if impact["equipment"] != expected_equipment:
            raise ReportKeyScenariosError(
                f"Результаты для сценария {code} устарели: "
                "не совпадает оборудование или составляющая ОПО"
            )
        zones = []
        for field, _ in ZONE_FIELDS:
            value = impact[field]
            if value != "—":
                unit = "м²" if field == "spill_area_m2" else "м"
                zones.append(f"{ZONE_DESCRIPTIONS[field]} — {value} {unit}")
        rows.append(
            {
                "component": str(item["hazard_component"]),
                "scenario_type": str(item["scenario_type_name"]),
                "scenario_code": code,
                "factor": impact["factor"],
                "zones": "; ".join(zones) or "Зоны поражения отсутствуют",
            }
        )
    return tuple(rows)


def load_key_scenario_people_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    try:
        selected = KeyScenariosService().calculate(project_directory).rows
    except KeyScenariosError as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    return tuple(
        {
            "component": str(item["hazard_component"]),
            "scenario_type": str(item["scenario_type_name"]),
            "scenario_code": str(item["scenario_code"]),
            "fatalities": str(item["fatalities_count"]),
            "injured": str(item["injured_count"]),
            "affected": str(item["fatalities_count"] + item["injured_count"]),
        }
        for item in selected
    )


def load_key_scenario_damage_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    try:
        selected = KeyScenariosService().calculate(project_directory).rows
        damage_rows = load_damage_rows(project_directory)
    except (KeyScenariosError, ReportDamageError) as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    damage_by_code = {item["code"]: item for item in damage_rows}
    rows: list[dict[str, str]] = []
    for item in selected:
        code = str(item["scenario_code"])
        damage = damage_by_code.get(code)
        if damage is None:
            raise ReportKeyScenariosError(
                f"В damage_results.json отсутствует ключевой сценарий {code}"
            )
        expected_equipment = (
            f"{item['equipment_name']} ({item['hazard_component']})"
        )
        if damage["equipment"] != expected_equipment:
            raise ReportKeyScenariosError(
                f"Результаты для сценария {code} устарели: "
                "не совпадает оборудование или составляющая ОПО"
            )
        row = {
            "component": str(item["hazard_component"]),
            "scenario_type": str(item["scenario_type_name"]),
            "scenario_code": code,
        }
        row.update({field: damage[field] for field, _ in DAMAGE_FIELDS})
        rows.append(row)
    return tuple(rows)


def load_key_scenario_conclusions(
    project_directory: Path | str,
) -> tuple[str, ...]:
    try:
        selected = KeyScenariosService().calculate(project_directory).rows
    except KeyScenariosError as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    grouped: dict[str, dict[str, Any]] = {}
    for item in selected:
        grouped.setdefault(str(item["hazard_component"]), {})[
            str(item["scenario_type"])
        ] = item

    conclusions: list[str] = []
    for component, values in grouped.items():
        dangerous = values.get("dangerous")
        probable = values.get("probable")
        if dangerous is None or probable is None:
            raise ReportKeyScenariosError(
                f"Для составляющей ОПО «{component}» не определены оба "
                "ключевых сценария"
            )
        dangerous_code = str(dangerous["scenario_code"])
        damage = f"{float(dangerous['total_damage']):.1f}".replace(".", ",")
        if dangerous_code == str(probable["scenario_code"]):
            frequency = f"{float(dangerous['scenario_frequency']):.3E}"
            conclusions.append(
                f"Для составляющей ОПО «{component}» наиболее опасным и "
                f"наиболее вероятным является сценарий {dangerous_code}: "
                f"погибло {dangerous['fatalities_count']} чел., ранено "
                f"{dangerous['injured_count']} чел., всего пострадало "
                f"{dangerous['fatalities_count'] + dangerous['injured_count']} чел., "
                "суммарный ущерб — "
                f"{damage} тыс. руб., частота — {frequency} 1/год."
            )
            continue
        frequency = f"{float(probable['scenario_frequency']):.3E}"
        conclusions.append(
            f"Для составляющей ОПО «{component}» наиболее опасным является "
            f"сценарий {dangerous_code}: погибло "
            f"{dangerous['fatalities_count']} чел., ранено "
            f"{dangerous['injured_count']} чел., всего пострадало "
            f"{dangerous['fatalities_count'] + dangerous['injured_count']} чел., "
            "суммарный ущерб — "
            f"{damage} тыс. руб. Наиболее вероятным является сценарий "
            f"{probable['scenario_code']} с частотой {frequency} 1/год."
        )
    return tuple(conclusions)


def load_accident_description_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    project = Path(project_directory)
    try:
        selected = KeyScenariosService().calculate(project).rows
        factors = load_key_scenario_pf_rows(project)
        masses = load_ov_amount_rows(project)
    except (KeyScenariosError, ReportOvAmountError) as exc:
        raise ReportKeyScenariosError(str(exc)) from exc

    try:
        raw = json.loads((project / IMPACT_ZONES_FILE_NAME).read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ReportKeyScenariosError(
            "Зоны поражающих факторов не рассчитаны. "
            "Сначала выполните модуль «Зоны ПФ»"
        ) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportKeyScenariosError(
            f"Не удалось прочитать {IMPACT_ZONES_FILE_NAME}"
        ) from exc
    impact_values = raw.get("results") if isinstance(raw, dict) else None
    if not isinstance(impact_values, list):
        raise ReportKeyScenariosError(
            f"Файл {IMPACT_ZONES_FILE_NAME} повреждён: results должен быть списком"
        )

    factor_by_code = {item["scenario_code"]: item for item in factors}
    mass_by_code = {item["code"]: item for item in masses}
    method_by_code: dict[str, str] = {}
    for item in impact_values:
        if not isinstance(item, dict):
            raise ReportKeyScenariosError(
                f"Файл {IMPACT_ZONES_FILE_NAME} содержит запись неверного формата"
            )
        code = str(item.get("scenario_code", "")).strip()
        calc_code = item.get("calc_code")
        if (
            not code
            or code in method_by_code
            or isinstance(calc_code, bool)
            or calc_code not in CALCULATION_METHODS
        ):
            raise ReportKeyScenariosError(
                f"Файл {IMPACT_ZONES_FILE_NAME} содержит некорректный сценарий"
            )
        method_by_code[code] = CALCULATION_METHODS[calc_code]

    rows: list[dict[str, str]] = []
    for item in selected:
        code = str(item["scenario_code"])
        factor = factor_by_code.get(code)
        mass = mass_by_code.get(code)
        method = method_by_code.get(code)
        if factor is None or mass is None or method is None:
            raise ReportKeyScenariosError(
                f"Для ключевого сценария {code} отсутствуют актуальные результаты расчёта"
            )
        fatalities = int(item["fatalities_count"])
        injured = int(item["injured_count"])
        rows.append(
            {
                "component": str(item["hazard_component"]),
                "scenario_type": str(item["scenario_type_name"]),
                "scenario_code": code,
                "equipment": str(item["equipment_name"]),
                "description": str(item["scenario_text"]),
                "frequency": f"{float(item['scenario_frequency']):.3E}",
                "accident_mass": mass["accident_mass"],
                "zones": factor["zones"],
                "method": method,
                "people": (
                    f"Пострадавшие: {fatalities + injured}\n"
                    f"Раненые: {injured}\nПогибшие: {fatalities}"
                ),
                "damage": f"{float(item['total_damage']):.1f}".replace(".", ","),
            }
        )
    return tuple(rows)


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _cell_margins(cell: Any, value: int = 50) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_cell_text(
    cell: Any,
    value: str,
    *,
    bold: bool = False,
    centered: bool = False,
    font_size: float = 8.5,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _cell_margins(cell)


def _paragraph_section(document: DocumentType, paragraph_element: Any) -> Any:
    section_index = 0
    for child in document.element.body:
        if child is paragraph_element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            section_index += 1
    return document.sections[min(section_index, len(document.sections) - 1)]


def _set_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.12, 0.11, 0.06, 0.17, 0.09, 0.09, 0.10, 0.14, 0.12)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _set_description_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.18, 0.16, 0.08, 0.58)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _set_pf_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.15, 0.14, 0.06, 0.18, 0.47)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _set_people_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.24, 0.20, 0.08, 0.15, 0.15, 0.18)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _set_damage_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.14, 0.11, 0.05, 0.11, 0.10, 0.12, 0.11, 0.14, 0.12)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _set_accident_description_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (
        0.11, 0.08, 0.045, 0.13, 0.15, 0.07,
        0.07, 0.14, 0.10, 0.065, 0.04,
    )
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        _prevent_split = row._tr.get_or_add_trPr()
        if _prevent_split.find(qn("w:cantSplit")) is None:
            _prevent_split.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def render_key_scenarios_section(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False

    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    for run in marker_paragraph.runs:
        run.text = ""
    marker_paragraph.paragraph_format.page_break_before = True
    marker_paragraph.paragraph_format.keep_with_next = True
    marker_paragraph.paragraph_format.space_before = Pt(0)
    marker_paragraph.paragraph_format.space_after = Pt(0)
    headers = (
        "Составляющая\nОПО",
        "Тип сценария",
        "№",
        "Оборудование",
        "Погибло,\nчел.",
        "Ранено,\nчел.",
        "Пострадало,\nчел.",
        "Суммарный ущерб,\nтыс. руб.",
        "Частота,\n1/год",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=7.5)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)

    for item in rows:
        cells = table.add_row().cells
        values = (
            item["component"],
            item["scenario_type"],
            item["scenario_code"],
            item["equipment"],
            item["fatalities"],
            item["injured"],
            item["affected"],
            item["damage"],
            item["frequency"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column not in (0, 1, 3))

    _set_table_geometry(section, table)
    return True


def render_key_scenario_descriptions(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if DESCRIPTION_MARKER in paragraph.text
        ),
        None,
    )
    if marker_paragraph is None:
        return False

    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Тип сценария",
        "№",
        "Описание сценария",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=8.5)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)

    for item in rows:
        cells = table.add_row().cells
        values = (
            item["component"],
            item["scenario_type"],
            item["scenario_code"],
            item["description"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column == 2, font_size=9)

    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_description_table_geometry(section, table)
    return True


def render_key_scenario_hazard_factors(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if PF_MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False

    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Тип сценария",
        "№",
        "Поражающий фактор",
        "Параметры зон поражения",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=8)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)

    for item in rows:
        cells = table.add_row().cells
        values = (
            item["component"],
            item["scenario_type"],
            item["scenario_code"],
            item["factor"],
            item["zones"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column == 2, font_size=8.5)

    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_pf_table_geometry(section, table)
    return True


def render_key_scenario_people(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if PEOPLE_MARKER in paragraph.text
        ),
        None,
    )
    if marker_paragraph is None:
        return False

    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Тип сценария",
        "№",
        "Погибло, чел.",
        "Ранено, чел.",
        "Пострадало, чел.",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=8.5)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)

    for item in rows:
        cells = table.add_row().cells
        values = (
            item["component"],
            item["scenario_type"],
            item["scenario_code"],
            item["fatalities"],
            item["injured"],
            item["affected"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column >= 2, font_size=9)

    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_people_table_geometry(section, table)
    return True


def render_key_scenario_damage(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if DAMAGE_MARKER in paragraph.text
        ),
        None,
    )
    if marker_paragraph is None:
        return False

    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=3 + len(DAMAGE_FIELDS))
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Тип сценария",
        "№",
    ) + tuple(label for _, label in DAMAGE_FIELDS)
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=7)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)

    for item in rows:
        cells = table.add_row().cells
        values = (
            item["component"],
            item["scenario_type"],
            item["scenario_code"],
        ) + tuple(item[field] for field, _ in DAMAGE_FIELDS)
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value, centered=column >= 2, font_size=7.5)

    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_damage_table_geometry(section, table)
    return True


def render_key_scenario_conclusions(
    document: DocumentType,
    conclusions: tuple[str, ...],
) -> bool:
    marker_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if CONCLUSION_MARKER in paragraph.text
        ),
        None,
    )
    if marker_paragraph is None:
        return False

    anchor = marker_paragraph._p
    for value in conclusions:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(value)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for name in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn(f"w:{name}"), "Times New Roman")
        anchor.addnext(paragraph._p)
        anchor = paragraph._p
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    return True


def render_accident_description_table(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if ACCIDENT_DESCRIPTION_MARKER in paragraph.text
        ),
        None,
    )
    if marker_paragraph is None:
        return False

    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=11)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    headers = (
        "Составляющая объекта",
        "Тип сценария",
        "Номер сценария",
        "Наименование оборудования",
        "Краткое описание сценария",
        "Частота, 1/год",
        "Количество ОВ, участвующего в аварии, т",
        "Зоны действия поражающих факторов",
        "Метод расчёта",
        "Количество людей",
        "Ущерб, тыс. руб.",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True, font_size=6.5)
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)

    for item in rows:
        values = (
            item["component"],
            item["scenario_type"],
            item["scenario_code"],
            item["equipment"],
            item["description"],
            item["frequency"],
            item["accident_mass"],
            item["zones"],
            item["method"],
            item["people"],
            item["damage"],
        )
        for column, (cell, value) in enumerate(zip(table.add_row().cells, values)):
            _set_cell_text(
                cell,
                value,
                centered=column in (2, 5, 6, 9, 10),
                font_size=6.5,
            )

    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_accident_description_table_geometry(section, table)
    return True


__all__ = [
    "ACCIDENT_DESCRIPTION_MARKER",
    "CONCLUSION_MARKER",
    "DAMAGE_MARKER",
    "DESCRIPTION_MARKER",
    "MARKER",
    "PF_MARKER",
    "PEOPLE_MARKER",
    "ReportKeyScenariosError",
    "load_key_scenario_conclusions",
    "load_accident_description_rows",
    "load_key_scenario_damage_rows",
    "load_key_scenario_description_rows",
    "load_key_scenario_people_rows",
    "load_key_scenario_pf_rows",
    "load_key_scenario_rows",
    "render_key_scenario_conclusions",
    "render_accident_description_table",
    "render_key_scenario_damage",
    "render_key_scenario_descriptions",
    "render_key_scenario_hazard_factors",
    "render_key_scenario_people",
    "render_key_scenarios_section",
]
