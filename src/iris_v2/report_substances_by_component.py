import math
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.equipment import PIPELINE_TYPES
from iris_v2.report_distribution import (
    ReportDistributionError,
    load_amount_results,
)
from iris_v2.report_equipment import ReportEquipmentError, load_project_equipment
from iris_v2.substances import SubstanceError, SubstanceService


MARKER = "{{SUBSTANCES_BY_COMPONENT_TABLE}}"
TOXIC_KINDS = {1, 3, 5, 6, 7, 8}
FLAMMABLE_GAS_KINDS = {2, 3, 4, 5}
FLAMMABLE_LIQUID_KINDS = {0, 1, 9}
STORAGE_WORDS = ("склад", "база", "хранилищ")
IDENTIFICATION_FIELDS = (
    "individual",
    "flammable_gas",
    "flammable_liquid_storage",
    "flammable_liquid_process",
    "toxic",
    "highly_toxic",
    "oxidizing",
    "explosive",
    "environmental",
)


class ReportSubstancesByComponentError(Exception):
    pass


def _mass(value: Any, label: str) -> float:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or float(value) < 0
    ):
        raise ReportSubstancesByComponentError(
            f"{label} должно быть числом не меньше нуля"
        )
    return float(value)


def _units(item: dict[str, Any], equipment_id: int) -> int:
    if item.get("equipment_type") in PIPELINE_TYPES:
        return 1
    value = item.get("equipment_count")
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or float(value) < 1
        or not float(value).is_integer()
    ):
        raise ReportSubstancesByComponentError(
            f"Оборудование {equipment_id}: equipment_count должно быть "
            "целым числом не меньше 1"
        )
    return int(value)


def _identification_flags(item: dict[str, Any]) -> dict[str, bool]:
    kind = item.get("kind")
    toxicity = item.get("toxicity")
    hazard_class = toxicity.get("hazard_class") if isinstance(toxicity, dict) else None
    flags = {
        "individual": False,
        "flammable_gas": kind in FLAMMABLE_GAS_KINDS,
        "flammable_liquid_storage": False,
        "flammable_liquid_process": kind in FLAMMABLE_LIQUID_KINDS,
        "toxic": kind in TOXIC_KINDS,
        "highly_toxic": hazard_class == 1 and kind in TOXIC_KINDS,
        "oxidizing": False,
        "explosive": False,
        "environmental": False,
    }
    configured = item.get("identification")
    if isinstance(configured, dict):
        for field in IDENTIFICATION_FIELDS:
            value = configured.get(field)
            if isinstance(value, bool):
                flags[field] = value
    return flags


def _format_mass(value: float) -> str:
    return f"{value:.6f}".rstrip("0").rstrip(".").replace(".", ",")


def load_substances_by_component_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    project = Path(project_directory)
    try:
        substances = SubstanceService().load_project(project)
        equipment = load_project_equipment(project)
        amounts = load_amount_results(project)
    except (SubstanceError, ReportEquipmentError, ReportDistributionError) as exc:
        raise ReportSubstancesByComponentError(str(exc)) from exc
    if not substances:
        raise ReportSubstancesByComponentError(
            "Вещества проекта не выбраны. Сначала заполните модуль «Вещества»"
        )

    substances_by_id: dict[int, dict[str, Any]] = {}
    for index, item in enumerate(substances, start=1):
        substance_id = item.get("id")
        name = str(item.get("name", "")).strip()
        if (
            isinstance(substance_id, bool)
            or not isinstance(substance_id, int)
            or substance_id <= 0
            or substance_id in substances_by_id
        ):
            raise ReportSubstancesByComponentError(
                f"substances.json, запись {index}: недопустимый или повторяющийся id"
            )
        if not name:
            raise ReportSubstancesByComponentError(
                f"Вещество {substance_id}: наименование не заполнено"
            )
        substances_by_id[substance_id] = item

    amount_by_equipment: dict[int, float] = {}
    for index, item in enumerate(amounts, start=1):
        equipment_id = item.get("equipment_id")
        if (
            isinstance(equipment_id, bool)
            or not isinstance(equipment_id, int)
            or equipment_id <= 0
            or equipment_id in amount_by_equipment
        ):
            raise ReportSubstancesByComponentError(
                f"amount_results.json, запись {index}: "
                "недопустимый или повторяющийся equipment_id"
            )
        amount_by_equipment[equipment_id] = _mass(
            item.get("amount_t"),
            f"amount_results.json, оборудование {equipment_id}: amount_t",
        )

    grouped: dict[int, dict[str, float]] = {}
    equipment_ids: set[int] = set()
    for index, item in enumerate(equipment, start=1):
        equipment_id = item.get("id")
        if (
            isinstance(equipment_id, bool)
            or not isinstance(equipment_id, int)
            or equipment_id <= 0
            or equipment_id in equipment_ids
        ):
            raise ReportSubstancesByComponentError(
                f"equipments.json, запись {index}: недопустимый или повторяющийся id"
            )
        equipment_ids.add(equipment_id)
        component = str(item.get("hazard_component", "")).strip()
        if not component:
            raise ReportSubstancesByComponentError(
                f"Оборудование {equipment_id}: составляющая ОПО не заполнена"
            )
        substance_id = item.get("substance_id")
        if (
            isinstance(substance_id, bool)
            or not isinstance(substance_id, int)
            or substance_id not in substances_by_id
        ):
            raise ReportSubstancesByComponentError(
                f"Оборудование {equipment_id}: substance_id отсутствует "
                "в substances.json"
            )
        if equipment_id not in amount_by_equipment:
            raise ReportSubstancesByComponentError(
                f"В amount_results.json отсутствует оборудование id={equipment_id}"
            )
        mass = amount_by_equipment[equipment_id] * _units(item, equipment_id)
        values = grouped.setdefault(
            substance_id,
            {field: 0.0 for field in ("total", *IDENTIFICATION_FIELDS)},
        )
        values["total"] += mass
        flags = _identification_flags(substances_by_id[substance_id])
        if flags["flammable_liquid_process"]:
            is_storage = any(word in component.lower() for word in STORAGE_WORDS)
            flags["flammable_liquid_process"] = not is_storage
            flags["flammable_liquid_storage"] = is_storage
        for field in IDENTIFICATION_FIELDS:
            if flags[field]:
                values[field] += mass

    extra_ids = set(amount_by_equipment) - equipment_ids
    if extra_ids:
        values = ", ".join(str(value) for value in sorted(extra_ids))
        raise ReportSubstancesByComponentError(
            "В amount_results.json найдено отсутствующее в equipments.json "
            f"оборудование: {values}"
        )
    return tuple(
        {
            "substance": str(substances_by_id[substance_id]["name"]).strip(),
            **{field: _format_mass(value) for field, value in values.items()},
        }
        for substance_id, values in grouped.items()
    )


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_text(
    cell: Any,
    value: str,
    *,
    bold: bool = False,
    centered: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    for existing_run in paragraph.runs:
        paragraph._p.remove(existing_run._r)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_text_direction(cell: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    direction = properties.find(qn("w:textDirection"))
    if direction is None:
        direction = OxmlElement("w:textDirection")
        properties.append(direction)
    direction.set(qn("w:val"), "btLr")


def _set_row_height(row: Any, value: int) -> None:
    properties = row._tr.get_or_add_trPr()
    height = properties.find(qn("w:trHeight"))
    if height is None:
        height = OxmlElement("w:trHeight")
        properties.append(height)
    height.set(qn("w:val"), str(value))
    height.set(qn("w:hRule"), "atLeast")


def _paragraph_section(document: DocumentType, paragraph_element: Any) -> Any:
    section_index = 0
    for child in document.element.body:
        if child is paragraph_element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            section_index += 1
    return document.sections[min(section_index, len(document.sections) - 1)]


def _set_table_geometry(section: Any, table: Any, *, landscape: bool = False) -> None:
    page_width = section.page_height if landscape else section.page_width
    horizontal_margins = 914400 if landscape else (
        section.left_margin + section.right_margin
    )
    total_twips = int(
        (page_width - horizontal_margins) / 635
    )
    proportions = (
        0.27, 0.15, 0.055, 0.08, 0.085, 0.085,
        0.05, 0.055, 0.05, 0.05, 0.07,
    )
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        column_index = 0
        for cell in row._tr.tc_lst:
            span_element = cell.get_or_add_tcPr().find(qn("w:gridSpan"))
            span = int(span_element.get(qn("w:val"))) if span_element is not None else 1
            width = sum(widths[column_index : column_index + span])
            column_index += span
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _section_break(properties: Any) -> Any:
    paragraph = OxmlElement("w:p")
    paragraph_properties = OxmlElement("w:pPr")
    section_properties = deepcopy(properties)
    section_type = section_properties.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        section_properties.insert(0, section_type)
    section_type.set(qn("w:val"), "nextPage")
    paragraph_properties.append(section_properties)
    paragraph.append(paragraph_properties)
    return paragraph


def _landscape_properties(section: Any) -> Any:
    properties = deepcopy(section._sectPr)
    page_size = properties.find(qn("w:pgSz"))
    width = page_size.get(qn("w:w"))
    height = page_size.get(qn("w:h"))
    page_size.set(qn("w:w"), height)
    page_size.set(qn("w:h"), width)
    page_size.set(qn("w:orient"), "landscape")
    margins = properties.find(qn("w:pgMar"))
    margins.set(qn("w:left"), "720")
    margins.set(qn("w:right"), "720")
    return properties


def render_substances_by_component_table(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False
    section = _paragraph_section(document, marker_paragraph._p)
    landscape = section.orientation != WD_ORIENT.LANDSCAPE
    caption_paragraph = marker_paragraph._p.getprevious()
    if landscape and caption_paragraph is not None:
        caption_paragraph.addprevious(_section_break(section._sectPr))
    table = document.add_table(rows=3, cols=11)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    _set_cell_text(table.cell(0, 0).merge(table.cell(0, 1)), "Вещество", bold=True, centered=True)
    _set_cell_text(
        table.cell(0, 2).merge(table.cell(0, 10)),
        "Признаки идентификации",
        bold=True,
        centered=True,
    )
    vertical_headers = {
        0: "Наименование",
        1: "Количество, т (всего)",
        2: "Индивидуально опасное вещество, т",
        3: "Воспламеняющиеся газы, т",
        6: "Токсичные вещества, т",
        7: "Высокотоксичные вещества, т",
        8: "Окисляющие вещества, т",
        9: "Взрывчатые вещества, т",
        10: "Вещества, опасные для окружающей среды, т",
    }
    for column, value in vertical_headers.items():
        cell = table.cell(1, column).merge(table.cell(2, column))
        _set_cell_text(cell, value, bold=True, centered=True)
        _set_text_direction(cell)
    _set_cell_text(
        table.cell(1, 4).merge(table.cell(1, 5)),
        "Горючие жидкости",
        bold=True,
        centered=True,
    )
    for column, value in (
        (4, "На складах и базах, т"),
        (5, "В технологическом процессе, т"),
    ):
        _set_cell_text(table.cell(2, column), value, bold=True, centered=True)
        _set_text_direction(table.cell(2, column))
    for header_row in table.rows[:3]:
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        header_row._tr.get_or_add_trPr().append(repeat_header)
    _set_row_height(table.rows[1], 2100)
    _set_row_height(table.rows[2], 1800)
    for item in rows:
        cells = table.add_row().cells
        values = (
            item["substance"], item["total"], item["individual"],
            item["flammable_gas"], item["flammable_liquid_storage"],
            item["flammable_liquid_process"], item["toxic"],
            item["highly_toxic"], item["oxidizing"], item["explosive"],
            item["environmental"],
        )
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(cell, value if value != "0" else "–", centered=True)
    totals = {
        field: sum(float(item[field].replace(",", ".")) for item in rows)
        for field in ("total", *IDENTIFICATION_FIELDS)
    }
    cells = table.add_row().cells
    _set_cell_text(cells[0], "Всего на ОПО:", bold=True, centered=True)
    for cell, field in zip(cells[1:], ("total", *IDENTIFICATION_FIELDS)):
        value = _format_mass(totals[field])
        _set_cell_text(cell, value if value != "0" else "–", bold=True, centered=True)
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    if landscape:
        table._tbl.addnext(_section_break(_landscape_properties(section)))
    _set_table_geometry(section, table, landscape=landscape)
    return True


__all__ = [
    "MARKER",
    "ReportSubstancesByComponentError",
    "load_substances_by_component_rows",
    "render_substances_by_component_table",
]
