import json
import math
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.calculation_config import (
    CalculationConfigError,
    CalculationConfigService,
)
from iris_v2.equipment import PIPELINE_TYPES
from iris_v2.report_distribution import (
    ReportDistributionError,
    load_amount_results,
)
from iris_v2.report_equipment import ReportEquipmentError, load_project_equipment
from iris_v2.substances import SubstanceError, SubstanceService


MARKER = "{{COMPONENT_INPUTS_ASSUMPTIONS_SECTION}}"


class ReportComponentInputsAssumptionsError(Exception):
    pass


def _number(value: Any, label: str, *, positive: bool = False) -> float:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or (positive and float(value) <= 0)
        or (not positive and float(value) < 0)
    ):
        condition = "больше нуля" if positive else "не меньше нуля"
        raise ReportComponentInputsAssumptionsError(
            f"{label} должно быть числом {condition}"
        )
    return float(value)


def _format(value: float, digits: int = 3) -> str:
    text = f"{value:.{digits}f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def _load_cases(project: Path) -> tuple[dict[str, Any], ...]:
    path = project / "calculation_cases.json"
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ReportComponentInputsAssumptionsError(
            "Расчётные сценарии не сформированы"
        ) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportComponentInputsAssumptionsError(
            "Не удалось прочитать calculation_cases.json"
        ) from exc
    values = raw.get("cases") if isinstance(raw, dict) else None
    if not isinstance(values, list) or not values:
        raise ReportComponentInputsAssumptionsError(
            "calculation_cases.json не содержит сценариев"
        )
    if any(not isinstance(item, dict) for item in values):
        raise ReportComponentInputsAssumptionsError(
            "calculation_cases.json содержит запись неверного формата"
        )
    return tuple(values)


def _units(item: dict[str, Any], equipment_id: int) -> int:
    if item.get("equipment_type") in PIPELINE_TYPES:
        return 1
    value = item.get("equipment_count")
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or float(value) < 1
        or not float(value).is_integer()
    ):
        raise ReportComponentInputsAssumptionsError(
            f"Оборудование {equipment_id}: equipment_count должно быть "
            "целым числом не меньше 1"
        )
    return int(value)


def _equipment_assumptions(item: dict[str, Any], equipment_id: int) -> str:
    values: list[str] = []
    for key, label, unit in (
        ("shutdown_time_s", "отключение", "с"),
        ("evaporation_time_s", "испарение", "с"),
        ("spill_coefficient", "растекание", "м⁻¹"),
        ("spill_area_m2", "площадь пролива", "м²"),
    ):
        raw = item.get(key)
        if raw is not None:
            number = _number(raw, f"Оборудование {equipment_id}: {key}")
            values.append(f"{label}: {_format(number)} {unit}")
    clutter = item.get("clutter_degree")
    if clutter is not None:
        if isinstance(clutter, bool) or not isinstance(clutter, int) or clutter not in range(1, 5):
            raise ReportComponentInputsAssumptionsError(
                f"Оборудование {equipment_id}: clutter_degree должен быть от 1 до 4"
            )
        values.append(f"загромождённость: {clutter}")
    return "; ".join(values) or "—"


def load_component_inputs_assumptions(
    project_directory: Path | str,
) -> tuple[tuple[dict[str, str], ...], tuple[tuple[str, str], ...]]:
    project = Path(project_directory)
    try:
        equipment = load_project_equipment(project)
        substances = SubstanceService().load_project(project)
        amounts = load_amount_results(project)
        config = CalculationConfigService().load(project)
    except (
        ReportEquipmentError,
        SubstanceError,
        ReportDistributionError,
        CalculationConfigError,
    ) as exc:
        raise ReportComponentInputsAssumptionsError(str(exc)) from exc
    cases = _load_cases(project)

    substance_names: dict[int, str] = {}
    for index, item in enumerate(substances, start=1):
        substance_id = item.get("id")
        name = str(item.get("name", "")).strip()
        if (
            isinstance(substance_id, bool)
            or not isinstance(substance_id, int)
            or substance_id <= 0
            or substance_id in substance_names
            or not name
        ):
            raise ReportComponentInputsAssumptionsError(
                f"substances.json, запись {index}: неверный id или наименование"
            )
        substance_names[substance_id] = name

    amount_by_id: dict[int, float] = {}
    for index, item in enumerate(amounts, start=1):
        equipment_id = item.get("equipment_id")
        if (
            isinstance(equipment_id, bool)
            or not isinstance(equipment_id, int)
            or equipment_id <= 0
            or equipment_id in amount_by_id
        ):
            raise ReportComponentInputsAssumptionsError(
                f"amount_results.json, запись {index}: неверный equipment_id"
            )
        amount_by_id[equipment_id] = _number(
            item.get("amount_t"),
            f"Оборудование {equipment_id}: amount_t",
        )

    scenarios_by_id: dict[int, list[str]] = {}
    for index, case in enumerate(cases, start=1):
        equipment_id = case.get("equipment_id")
        code = str(case.get("scenario_code", "")).strip()
        description = str(case.get("scenario_text", "")).strip()
        if (
            isinstance(equipment_id, bool)
            or not isinstance(equipment_id, int)
            or equipment_id <= 0
            or not code
            or not description
        ):
            raise ReportComponentInputsAssumptionsError(
                f"calculation_cases.json, запись {index}: неверные данные сценария"
            )
        scenarios_by_id.setdefault(equipment_id, []).append(
            f"{code}: {description}"
        )

    rows: list[dict[str, str]] = []
    equipment_ids: set[int] = set()
    for index, item in enumerate(equipment, start=1):
        equipment_id = item.get("id")
        if (
            isinstance(equipment_id, bool)
            or not isinstance(equipment_id, int)
            or equipment_id <= 0
            or equipment_id in equipment_ids
        ):
            raise ReportComponentInputsAssumptionsError(
                f"equipments.json, запись {index}: неверный или повторяющийся id"
            )
        equipment_ids.add(equipment_id)
        component = str(item.get("hazard_component", "")).strip()
        name = str(item.get("equipment_name", "")).strip()
        substance_id = item.get("substance_id")
        if not component or not name or substance_id not in substance_names:
            raise ReportComponentInputsAssumptionsError(
                f"Оборудование {equipment_id}: не заполнены основные данные"
            )
        if equipment_id not in amount_by_id:
            raise ReportComponentInputsAssumptionsError(
                f"В amount_results.json отсутствует оборудование id={equipment_id}"
            )
        phase = str(item.get("phase_state", "")).strip() or "—"
        pressure = _number(
            item.get("pressure_mpa"),
            f"Оборудование {equipment_id}: pressure_mpa",
        )
        temperature = _number(
            item.get("substance_temperature_c"),
            f"Оборудование {equipment_id}: substance_temperature_c",
        )
        total_mass = amount_by_id[equipment_id] * _units(item, equipment_id)
        rows.append(
            {
                "component": component,
                "equipment": name,
                "substance_mass": (
                    f"{substance_names[substance_id]}; {_format(total_mass)} т"
                ),
                "regime": (
                    f"{phase}; P={_format(pressure)} МПа; "
                    f"T={_format(temperature)} °C"
                ),
                "scenarios": "\n".join(scenarios_by_id.get(equipment_id, ()))
                or "Не сформированы",
                "assumptions": _equipment_assumptions(item, equipment_id),
            }
        )

    extra_amounts = set(amount_by_id) - equipment_ids
    extra_cases = set(scenarios_by_id) - equipment_ids
    if extra_amounts or extra_cases:
        raise ReportComponentInputsAssumptionsError(
            "Расчётные файлы содержат отсутствующее оборудование"
        )

    multipliers = config["frequency_multipliers"]
    assumptions = (
        ("Доля частичной разгерметизации", _format(config["partial_release_fraction"])),
        ("Доля вещества во взрывоопасном облаке", _format(config["flammable_cloud_fraction"])),
        ("Доля вещества в огненном шаре", _format(config["bleve_fraction"])),
        ("Доля частичного пролива", _format(config["partial_spill_fraction"])),
        ("Скорость ветра", f"{_format(config['wind_speed_m_s'])} м/с"),
        ("Коэффициент испарения", _format(config["evaporation_coefficient"])),
        ("Диаметр отверстия истечения жидкости", f"{_format(config['liquid_leak_hole_diameter_mm'])} мм"),
        ("Диаметр отверстия истечения газа", f"{_format(config['gas_leak_hole_diameter_mm'])} мм"),
        ("Плотность излучения огненного шара", f"{_format(config['fireball_surface_emissive_power_kw_m2'])} кВт/м²"),
        ("Масштаб ущерба", _format(config["damage_scale"])),
        ("Множители частот: стандартный / без КМ / с КМ", (
            f"{_format(multipliers['standard'])} / "
            f"{_format(multipliers['without_compensation'])} / "
            f"{_format(multipliers['with_compensation'])}"
        )),
    )
    return tuple(rows), assumptions


def _cell_text(cell: Any, value: str, *, bold: bool, centered: bool) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(7.5)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _repeat_header(row: Any) -> None:
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(repeat)


def _geometry(section: Any, table: Any, ratios: tuple[float, ...]) -> None:
    total = int((section.page_width - section.left_margin - section.right_margin) / 635)
    widths = [int(total * ratio) for ratio in ratios[:-1]]
    widths.append(total - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _insert_heading(document: DocumentType, marker: Any, text: str) -> Any:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    marker.addnext(paragraph._p)
    return paragraph._p


def render_component_inputs_assumptions_section(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
    assumptions: tuple[tuple[str, str], ...],
) -> bool:
    marker = next((p for p in document.paragraphs if MARKER in p.text), None)
    if marker is None:
        return False
    section = document.sections[-1]
    anchor = _insert_heading(
        document,
        marker._p,
        "Исходные данные по составляющим ОПО",
    )
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    anchor.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Оборудование",
        "Вещество; масса ОВ",
        "Состояние и режим",
        "Расчётные сценарии",
        "Локальные допущения",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _cell_text(cell, value, bold=True, centered=True)
    _repeat_header(table.rows[0])
    fields = (
        "component", "equipment", "substance_mass", "regime",
        "scenarios", "assumptions",
    )
    for item in rows:
        for index, (cell, field) in enumerate(zip(table.add_row().cells, fields)):
            _cell_text(cell, item[field], bold=False, centered=index in (2, 3))
    _geometry(section, table, (0.16, 0.21, 0.13, 0.14, 0.21, 0.15))

    assumptions_heading = _insert_heading(
        document,
        table._tbl,
        "Общие расчётные допущения",
    )
    assumptions_table = document.add_table(rows=1, cols=2)
    assumptions_table.style = "Table Grid"
    assumptions_heading.addnext(assumptions_table._tbl)
    for cell, value in zip(
        assumptions_table.rows[0].cells,
        ("Параметр", "Принятое значение"),
    ):
        _cell_text(cell, value, bold=True, centered=True)
    _repeat_header(assumptions_table.rows[0])
    for name, value in assumptions:
        cells = assumptions_table.add_row().cells
        _cell_text(cells[0], name, bold=False, centered=False)
        _cell_text(cells[1], value, bold=False, centered=True)
    _geometry(section, assumptions_table, (0.68, 0.32))
    marker._element.getparent().remove(marker._element)
    return True


__all__ = [
    "MARKER",
    "ReportComponentInputsAssumptionsError",
    "load_component_inputs_assumptions",
    "render_component_inputs_assumptions_section",
]
