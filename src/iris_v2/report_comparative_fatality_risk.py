import math
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from iris_v2.report_component_fatality_risk import (
    NOT_CALCULATED,
    ReportComponentFatalityRiskError,
    load_component_fatality_risk_rows,
)


MARKER = "{{COMPARATIVE_FATALITY_RISK_TABLE}}"
BACKGROUND_RISK_PPM = 195.0
BACKGROUND_RISK_PER_YEAR = BACKGROUND_RISK_PPM / 1_000_000
SOURCE_NOTE = (
    "Примечание — RГЛ = 195 ppm по таблице № 3 приложения № 2 к приказу "
    "Ростехнадзора от 12.09.2023 № 331."
)


class ReportComparativeFatalityRiskError(Exception):
    pass


def _decimal(value: float, digits: int = 1) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def _comparison(ppm: float) -> str:
    if ppm == 0:
        return "Ниже фонового значения"
    if math.isclose(ppm, BACKGROUND_RISK_PPM, rel_tol=1e-12, abs_tol=0.0):
        return "Соответствует фоновому значению"
    if ppm < BACKGROUND_RISK_PPM:
        return f"Ниже в {_decimal(BACKGROUND_RISK_PPM / ppm)} раза"
    return f"Выше в {_decimal(ppm / BACKGROUND_RISK_PPM)} раза"


def load_comparative_fatality_risk_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    try:
        source_rows = load_component_fatality_risk_rows(project_directory)
    except ReportComponentFatalityRiskError as exc:
        raise ReportComparativeFatalityRiskError(str(exc)) from exc

    rows: list[dict[str, str]] = []
    for item in source_rows:
        individual_text = item["individual"]
        if individual_text == NOT_CALCULATED:
            rows.append(
                {
                    "component": item["component"],
                    "risk": NOT_CALCULATED,
                    "ppm": "—",
                    "dbr": "—",
                    "comparison": NOT_CALCULATED,
                }
            )
            continue
        try:
            risk = float(individual_text)
        except ValueError as exc:
            raise ReportComparativeFatalityRiskError(
                "Не удалось преобразовать индивидуальный риск гибели"
            ) from exc
        ppm = risk * 1_000_000
        dbr = float("-inf") if ppm == 0 else 10 * math.log10(
            ppm / BACKGROUND_RISK_PPM
        )
        rows.append(
            {
                "component": item["component"],
                "risk": f"{risk:.3E}",
                "ppm": _decimal(ppm, 3),
                "dbr": "−∞" if not math.isfinite(dbr) else _decimal(dbr),
                "comparison": _comparison(ppm),
            }
        )
    rows.append(
        {
            "component": "Фоновый риск RГЛ",
            "risk": f"{BACKGROUND_RISK_PER_YEAR:.3E}",
            "ppm": _decimal(BACKGROUND_RISK_PPM, 0),
            "dbr": "0,0",
            "comparison": "Фоновое значение",
        }
    )
    return tuple(rows)


def _shade(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_text(
    cell: Any,
    value: str,
    *,
    bold: bool = False,
    centered: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_note_text(paragraph: Any) -> None:
    paragraph.text = ""
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(SOURCE_NOTE)
    run.font.name = "Times New Roman"
    run.font.size = Pt(7.5)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")


def _paragraph_section(document: DocumentType, paragraph_element: Any) -> Any:
    section_index = 0
    for child in document.element.body:
        if child is paragraph_element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            section_index += 1
    return document.sections[min(section_index, len(document.sections) - 1)]


def _set_table_geometry(section: Any, table: Any) -> None:
    total_twips = int(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    proportions = (0.27, 0.17, 0.13, 0.13, 0.30)
    widths = [int(total_twips * value) for value in proportions[:-1]]
    widths.append(total_twips - sum(widths))
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def render_comparative_fatality_risk_table(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker_paragraph = next(
        (paragraph for paragraph in document.paragraphs if MARKER in paragraph.text),
        None,
    )
    if marker_paragraph is None:
        return False
    section = _paragraph_section(document, marker_paragraph._p)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    marker_paragraph._p.addnext(table._tbl)
    note = document.add_paragraph()
    table._tbl.addnext(note._p)
    _set_note_text(note)
    headers = (
        "Составляющая ОПО / показатель",
        "Индивидуальный риск, 1/год",
        "Риск, ppm",
        "Риск, дБR",
        "Сравнение с RГЛ",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell_text(cell, value, bold=True, centered=True)
        _shade(cell, "D9E1F2")
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for row_index, item in enumerate(rows):
        cells = table.add_row().cells
        values = (
            item["component"],
            item["risk"],
            item["ppm"],
            item["dbr"],
            item["comparison"],
        )
        is_background = row_index == len(rows) - 1
        for column, (cell, value) in enumerate(zip(cells, values)):
            _set_cell_text(
                cell,
                value,
                bold=is_background,
                centered=column in (1, 2, 3),
            )
            if is_background:
                _shade(cell, "E2F0D9")
    marker_paragraph._element.getparent().remove(marker_paragraph._element)
    _set_table_geometry(section, table)
    return True


__all__ = [
    "BACKGROUND_RISK_PER_YEAR",
    "BACKGROUND_RISK_PPM",
    "MARKER",
    "ReportComparativeFatalityRiskError",
    "load_comparative_fatality_risk_rows",
    "render_comparative_fatality_risk_table",
]
