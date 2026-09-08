import json
import math
from pathlib import Path
from typing import Any

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


MARKER = "{{COMPONENT_RISK_SUMMARY_TABLE}}"
NOT_CALCULATED = "Не рассчитан"


class ReportComponentRiskSummaryError(Exception):
    pass


def _object(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ReportComponentRiskSummaryError(
            f"Файл не найден: {path.name}"
        ) from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportComponentRiskSummaryError(
            f"Не удалось прочитать {path.name}"
        ) from exc
    if not isinstance(value, dict):
        raise ReportComponentRiskSummaryError(
            f"Файл {path.name} повреждён: ожидается объект"
        )
    return value


def _number(value: Any, label: str) -> float:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or float(value) < 0
    ):
        raise ReportComponentRiskSummaryError(
            f"{label} должно быть числом не меньше нуля"
        )
    return float(value)


def _count(value: Any, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ReportComponentRiskSummaryError(
            f"{label} должно быть целым числом не меньше нуля"
        )
    return value


def _same(left: float, right: float) -> bool:
    return math.isclose(left, right, rel_tol=1e-12, abs_tol=1e-9)


def load_component_risk_summary_rows(
    project_directory: Path | str,
) -> tuple[dict[str, str], ...]:
    project = Path(project_directory)
    risk = _object(project / "risk_results.json")
    summary = _object(project / "risk_summary.json")
    results = risk.get("results")
    components = summary.get("components")
    if not isinstance(results, list) or not results:
        raise ReportComponentRiskSummaryError(
            "risk_results.json не содержит результатов"
        )
    if not isinstance(components, list):
        raise ReportComponentRiskSummaryError(
            "risk_summary.json не содержит составляющих ОПО"
        )

    grouped: dict[str, dict[str, Any]] = {}
    codes: set[str] = set()
    for index, item in enumerate(results, start=1):
        if not isinstance(item, dict):
            raise ReportComponentRiskSummaryError(
                f"risk_results.json, запись {index}: ожидается объект"
            )
        code = str(item.get("scenario_code", "")).strip()
        name = str(item.get("hazard_component", "")).strip()
        if not code or code in codes or not name:
            raise ReportComponentRiskSummaryError(
                f"risk_results.json, запись {index}: неверный сценарий или составляющая"
            )
        codes.add(code)
        fatalities = _count(item.get("fatalities_count"), f"Сценарий {code}: погибшие")
        injured = _count(item.get("injured_count"), f"Сценарий {code}: раненые")
        individual_raw = item.get("individual_risk_fatalities")
        individual = (
            None
            if individual_raw is None
            else _number(individual_raw, f"Сценарий {code}: индивидуальный риск")
        )
        values = grouped.setdefault(
            name,
            {
                "count": 0,
                "frequency": 0.0,
                "fatalities": 0,
                "injured": 0,
                "affected": 0,
                "damage": 0.0,
                "collective": 0.0,
                "individual": 0.0,
                "has_individual": False,
            },
        )
        values["count"] += 1
        values["frequency"] += _number(
            item.get("scenario_frequency"), f"Сценарий {code}: частота"
        )
        values["fatalities"] = max(values["fatalities"], fatalities)
        values["injured"] = max(values["injured"], injured)
        values["affected"] = max(values["affected"], fatalities + injured)
        values["damage"] = max(
            values["damage"],
            _number(item.get("total_damage"), f"Сценарий {code}: ущерб"),
        )
        values["collective"] += _number(
            item.get("collective_risk_fatalities"),
            f"Сценарий {code}: коллективный риск",
        )
        if individual is not None:
            values["individual"] += individual
            values["has_individual"] = True

    if (
        summary.get("case_count") != len(results)
        or summary.get("component_count") != len(grouped)
        or summary.get("risk_unit") != "1/год"
        or summary.get("damage_unit") != "тыс. руб."
        or len(components) != len(grouped)
    ):
        raise ReportComponentRiskSummaryError("risk_summary.json устарел")

    rows: list[dict[str, str]] = []
    for index, ((name, expected), stored) in enumerate(
        zip(grouped.items(), components), start=1
    ):
        if not isinstance(stored, dict) or stored.get("hazard_component") != name:
            raise ReportComponentRiskSummaryError(
                f"risk_summary.json, составляющая {index}: неверное наименование"
            )
        stored_individual = stored.get("individual_risk_fatalities")
        checks = (
            stored.get("scenario_count") == expected["count"],
            _same(_number(stored.get("max_total_damage"), name), expected["damage"]),
            _same(
                _number(stored.get("collective_risk_fatalities"), name),
                expected["collective"],
            ),
            (stored_individual is None) == (not expected["has_individual"]),
        )
        if not all(checks):
            raise ReportComponentRiskSummaryError(
                f"risk_summary.json устарел: составляющая «{name}»"
            )
        if stored_individual is not None and not _same(
            _number(stored_individual, name), expected["individual"]
        ):
            raise ReportComponentRiskSummaryError(
                f"risk_summary.json устарел: составляющая «{name}»"
            )
        rows.append(
            {
                "component": name,
                "frequency": f"{expected['frequency']:.3E}",
                "fatalities": str(expected["fatalities"]),
                "injured": str(expected["injured"]),
                "affected": str(expected["affected"]),
                "damage": f"{expected['damage']:.1f}".replace(".", ","),
                "collective": f"{expected['collective']:.3E}",
                "individual": (
                    f"{expected['individual']:.3E}"
                    if expected["has_individual"]
                    else NOT_CALCULATED
                ),
            }
        )
    return tuple(rows)


def _cell_text(cell: Any, value: str, *, bold: bool, centered: bool) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(7.5)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _section(document: DocumentType, element: Any) -> Any:
    index = 0
    for child in document.element.body:
        if child is element:
            break
        if child.tag == qn("w:p") and child.find(
            f"./{qn('w:pPr')}/{qn('w:sectPr')}"
        ) is not None:
            index += 1
    return document.sections[min(index, len(document.sections) - 1)]


def _geometry(section: Any, table: Any) -> None:
    total = int((section.page_width - section.left_margin - section.right_margin) / 635)
    ratios = (0.23, 0.12, 0.08, 0.08, 0.09, 0.14, 0.13)
    widths = [int(total * ratio) for ratio in ratios]
    widths.append(total - sum(widths))
    table.autofit = False
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table._tbl.tblPr.append(table_width)
    table_width.set(qn("w:w"), str(total))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row._tr.tc_lst, widths):
            cell_width = cell.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def render_component_risk_summary_table(
    document: DocumentType,
    rows: tuple[dict[str, str], ...],
) -> bool:
    marker = next((p for p in document.paragraphs if MARKER in p.text), None)
    if marker is None:
        return False
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    marker._p.addnext(table._tbl)
    headers = (
        "Составляющая ОПО",
        "Частота сценариев, 1/год",
        "Макс. погибших, чел.",
        "Макс. раненых, чел.",
        "Макс. пострадавших, чел.",
        "Макс. ущерб, тыс. руб.",
        "Коллективный риск гибели, чел./год",
        "Индивидуальный риск гибели, 1/год",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        _cell_text(cell, value, bold=True, centered=True)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    fields = (
        "component", "frequency", "fatalities", "injured", "affected",
        "damage", "collective", "individual",
    )
    for item in rows:
        for index, (cell, field) in enumerate(zip(table.add_row().cells, fields)):
            _cell_text(cell, item[field], bold=False, centered=index > 0)
    section = _section(document, marker._p)
    marker._element.getparent().remove(marker._element)
    _geometry(section, table)
    return True


__all__ = [
    "MARKER",
    "ReportComponentRiskSummaryError",
    "load_component_risk_summary_rows",
    "render_component_risk_summary_table",
]
