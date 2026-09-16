import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from iris_v2.report_event_trees import (
    EventTreeReportItem,
    MARKER,
    prepare_event_trees,
    render_event_trees_section,
)
from iris_v2.service import CreateProjectData, ProjectService


def _append_sequence_field(paragraph, value: int) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " SEQ Рисунок \\* ARABIC "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(str(value))
    paragraph.add_run()._r.append(end)


def _project(tmp_path: Path) -> Path:
    project = tmp_path / "project"
    ProjectService().create(
        project,
        CreateProjectData(
            name="Проект",
            code="P-1",
            organization_name="Организация",
            opo_name="ОПО",
            opo_registration_number="А00-00000-0000",
            organization_snapshot={},
            opo_snapshot={},
        ),
    )
    return project


def test_prepare_event_trees_removes_duplicate_pairs(tmp_path: Path) -> None:
    project = _project(tmp_path)
    cases = [
        {"scenario_code": "С1", "equipment_type": 0, "kind": 0},
        {"scenario_code": "С2", "equipment_type": 0, "kind": 0},
        {"scenario_code": "С3", "equipment_type": 1, "kind": 9},
    ]
    (project / "calculation_cases.json").write_text(
        json.dumps({"cases": cases}, ensure_ascii=False), encoding="utf-8"
    )

    items = prepare_event_trees(project)

    assert [(item.equipment_type, item.kind) for item in items] == [(0, 0), (1, 9)]
    assert all(item.path.is_file() for item in items)
    assert len(list((project / "output" / "event_trees").glob("*.png"))) == 2


def test_render_event_tree_replaces_marker(tmp_path: Path) -> None:
    image_path = (
        Path(__file__).parents[1]
        / "src"
        / "iris_v2"
        / "data"
        / "event_trees"
        / "event_tree_eq00_kind00.png"
    )
    document = Document()
    preceding_caption = document.add_paragraph("Рисунок ")
    _append_sequence_field(preceding_caption, 1)
    preceding_caption.add_run(" – Предыдущий рисунок")
    document.add_paragraph(MARKER)
    following_caption = document.add_paragraph("Рисунок ")
    _append_sequence_field(following_caption, 2)
    following_caption.add_run(" – Последующий рисунок")
    items = (
        EventTreeReportItem(0, 0, "Трубопровод", "ЛВЖ", image_path),
    )

    assert render_event_trees_section(document, items)

    text = "\n".join(p.text for p in document.paragraphs)
    assert MARKER not in text
    assert len(document.inline_shapes) == 1
    assert "Дерево событий представлено на рисунке 2." in text
    event_tree_caption = next(
        paragraph
        for paragraph in document.paragraphs
        if "Дерево событий для типа оборудования" in paragraph.text
    )
    assert event_tree_caption.text.startswith("Рисунок 2 –")
    assert following_caption.text.startswith("Рисунок 3 –")
    instructions = "".join(
        node.text or "" for node in document.element.body.iter(qn("w:instrText"))
    )
    assert "SEQ Рисунок" in instructions
    assert "REF IrisEventTreeFigure1" in instructions
    assert "ДС-" not in text
