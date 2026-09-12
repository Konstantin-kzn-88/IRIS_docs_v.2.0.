import json
from pathlib import Path

from docx import Document

from iris_v2.calculation_config import (
    CalculationConfigService,
    new_calculation_config,
)
from iris_v2.report_key_scenarios import (
    load_accident_description_rows,
    load_key_scenario_conclusions,
    load_key_scenario_damage_rows,
    load_key_scenario_description_rows,
    load_key_scenario_people_rows,
    load_key_scenario_pf_rows,
    load_key_scenario_rows,
    render_key_scenario_conclusions,
    render_key_scenario_damage,
    render_key_scenario_descriptions,
    render_key_scenario_hazard_factors,
    render_key_scenario_people,
    render_key_scenarios_section,
    render_accident_description_table,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def row(
    code: str,
    component: str,
    fatalities: int,
    injured: int,
    damage: float,
    frequency: float,
) -> dict:
    return {
        "scenario_code": code,
        "hazard_component": component,
        "equipment_name": f"Оборудование {code}",
        "fatalities_count": fatalities,
        "injured_count": injured,
        "total_damage": damage,
        "scenario_frequency": frequency,
        "scenario_text": f"Описание {code}",
    }


def test_rows_use_key_scenario_selection_and_report_formats(tmp_path: Path) -> None:
    write_json(
        tmp_path / "risk_results.json",
        {
            "results": [
                row("С1", "Участок", 1, 3, 1200.25, 4e-5),
                row("С2", "Участок", 2, 4, 900.0, 2e-5),
                row("С3", "Участок", 0, 1, 500.0, 6e-5),
            ]
        },
    )

    rows = load_key_scenario_rows(tmp_path)

    assert rows == (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "equipment": "Оборудование С2",
            "fatalities": "2",
            "injured": "4",
            "affected": "6",
            "damage": "900,0",
            "frequency": "2.000E-05",
        },
        {
            "component": "Участок",
            "scenario_type": "Наиболее вероятный",
            "scenario_code": "С3",
            "equipment": "Оборудование С3",
            "fatalities": "0",
            "injured": "1",
            "affected": "1",
            "damage": "500,0",
            "frequency": "6.000E-05",
        },
    )
    assert (tmp_path / "key_scenarios.json").is_file()


def test_section_replaces_marker_with_repeatable_table() -> None:
    document = Document()
    document.add_paragraph("{{TOP_SCENARIOS_BY_COMPONENT_SECTION}}")
    rows = (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С1",
            "equipment": "Аппарат",
            "fatalities": "1",
            "injured": "2",
            "affected": "3",
            "damage": "100,0",
            "frequency": "1.000E-05",
        },
    )

    assert render_key_scenarios_section(document, rows)
    assert "TOP_SCENARIOS" not in "\n".join(p.text for p in document.paragraphs)
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Участок",
        "Наиболее опасный",
        "С1",
        "Аппарат",
        "1",
        "2",
        "3",
        "100,0",
        "1.000E-05",
    ]
    properties = document.tables[0].rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None


def test_description_rows_use_selected_scenarios_and_full_text(
    tmp_path: Path,
) -> None:
    write_json(
        tmp_path / "risk_results.json",
        {
            "results": [
                row("С1", "Участок", 1, 3, 1200.25, 4e-5),
                row("С2", "Участок", 2, 4, 900.0, 2e-5),
                row("С3", "Участок", 0, 1, 500.0, 6e-5),
            ]
        },
    )

    assert load_key_scenario_description_rows(tmp_path) == (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "description": "Описание С2",
        },
        {
            "component": "Участок",
            "scenario_type": "Наиболее вероятный",
            "scenario_code": "С3",
            "description": "Описание С3",
        },
    )


def test_description_table_replaces_marker_and_repeats_header() -> None:
    document = Document()
    document.add_paragraph("{{TOP_SCENARIOS_DESC_BY_COMPONENT}}")
    rows = (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С1",
            "description": "Разрыв трубопровода → пожар пролива",
        },
    )

    assert render_key_scenario_descriptions(document, rows)
    assert "TOP_SCENARIOS_DESC" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Участок",
        "Наиболее опасный",
        "С1",
        "Разрыв трубопровода → пожар пролива",
    ]
    properties = document.tables[0].rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None


def test_pf_rows_use_selected_scenarios_and_calculated_zones(
    tmp_path: Path,
) -> None:
    risks = [
        row("С1", "Участок", 1, 3, 1200.25, 4e-5),
        row("С2", "Участок", 2, 4, 900.0, 2e-5),
        row("С3", "Участок", 0, 1, 500.0, 6e-5),
    ]
    write_json(tmp_path / "risk_results.json", {"results": risks})
    factors = []
    impacts = []
    for index, risk in enumerate(risks, start=1):
        calc_code = {1: 1, 2: 2, 3: 0}[index]
        common = {
            "id": index,
            "scenario_code": risk["scenario_code"],
            "equipment_name": risk["equipment_name"],
            "hazard_component": risk["hazard_component"],
            "calc_code": calc_code,
        }
        factors.append(common)
        impact = dict(common)
        impact["impact_values"] = {"p_2_m": 42.04} if index == 2 else {}
        impacts.append(impact)
    write_json(tmp_path / "hazard_factor_results.json", {"results": factors})
    write_json(tmp_path / "impact_zones.json", {"results": impacts})

    assert load_key_scenario_pf_rows(tmp_path) == (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "factor": "Воздушная ударная волна",
            "zones": "зона разрушения остекления (2 кПа) — 42,0 м",
        },
        {
            "component": "Участок",
            "scenario_type": "Наиболее вероятный",
            "scenario_code": "С3",
            "factor": "Поражающий фактор отсутствует",
            "zones": "Зоны поражения отсутствуют",
        },
    )


def test_pf_table_replaces_marker_and_repeats_header() -> None:
    document = Document()
    document.add_paragraph("{{TOP_SCENARIOS_PF_BY_COMPONENT}}")
    rows = (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "factor": "взрыв облака",
            "zones": "зона разрушения остекления (2 кПа) — 42,0 м",
        },
    )

    assert render_key_scenario_hazard_factors(document, rows)
    assert "TOP_SCENARIOS_PF" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Участок",
        "Наиболее опасный",
        "С2",
        "взрыв облака",
        "зона разрушения остекления (2 кПа) — 42,0 м",
    ]
    properties = document.tables[0].rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None


def test_people_rows_use_selected_scenarios(tmp_path: Path) -> None:
    write_json(
        tmp_path / "risk_results.json",
        {
            "results": [
                row("С1", "Участок", 1, 3, 1200.25, 4e-5),
                row("С2", "Участок", 2, 4, 900.0, 2e-5),
                row("С3", "Участок", 0, 1, 500.0, 6e-5),
            ]
        },
    )

    assert load_key_scenario_people_rows(tmp_path) == (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "fatalities": "2",
            "injured": "4",
            "affected": "6",
        },
        {
            "component": "Участок",
            "scenario_type": "Наиболее вероятный",
            "scenario_code": "С3",
            "fatalities": "0",
            "injured": "1",
            "affected": "1",
        },
    )


def test_people_table_replaces_marker_and_repeats_header() -> None:
    document = Document()
    document.add_paragraph("{{TOP_SCENARIOS_FATALITIES_INJURED}}")
    rows = (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "fatalities": "2",
            "injured": "4",
            "affected": "6",
        },
    )

    assert render_key_scenario_people(document, rows)
    assert "TOP_SCENARIOS_FATALITIES" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Участок",
        "Наиболее опасный",
        "С2",
        "2",
        "4",
        "6",
    ]
    properties = document.tables[0].rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None


def test_damage_rows_use_selected_scenarios_and_damage_breakdown(
    tmp_path: Path,
) -> None:
    risks = [
        row("С1", "Участок", 1, 3, 1200.0, 4e-5),
        row("С2", "Участок", 2, 4, 900.0, 2e-5),
        row("С3", "Участок", 0, 1, 500.0, 6e-5),
    ]
    write_json(tmp_path / "risk_results.json", {"results": risks})
    people = []
    damages = []
    for risk in risks:
        source = {
            "scenario_code": risk["scenario_code"],
            "equipment_name": risk["equipment_name"],
            "hazard_component": risk["hazard_component"],
        }
        people.append(source)
        damage = dict(source)
        damage.update(
            {
                "damage_scale": 30.3,
                "damage_unit": "тыс. руб.",
                "direct_losses": risk["total_damage"],
                "liquidation_costs": 0.0,
                "social_losses": 0.0,
                "indirect_damage": 0.0,
                "total_environmental_damage": 0.0,
                "total_damage": risk["total_damage"],
            }
        )
        damages.append(damage)
    write_json(tmp_path / "people_results.json", {"results": people})
    write_json(tmp_path / "damage_results.json", {"results": damages})
    CalculationConfigService().save(tmp_path, new_calculation_config())

    assert load_key_scenario_damage_rows(tmp_path) == (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "direct_losses": "900,0",
            "liquidation_costs": "0,0",
            "social_losses": "0,0",
            "indirect_damage": "0,0",
            "total_environmental_damage": "0,0",
            "total_damage": "900,0",
        },
        {
            "component": "Участок",
            "scenario_type": "Наиболее вероятный",
            "scenario_code": "С3",
            "direct_losses": "500,0",
            "liquidation_costs": "0,0",
            "social_losses": "0,0",
            "indirect_damage": "0,0",
            "total_environmental_damage": "0,0",
            "total_damage": "500,0",
        },
    )


def test_damage_table_replaces_marker_and_repeats_header() -> None:
    document = Document()
    document.add_paragraph("{{TOP_SCENARIOS_DAMAGE}}")
    rows = (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "direct_losses": "100,0",
            "liquidation_costs": "10,0",
            "social_losses": "20,0",
            "indirect_damage": "30,0",
            "total_environmental_damage": "40,0",
            "total_damage": "200,0",
        },
    )

    assert render_key_scenario_damage(document, rows)
    assert "TOP_SCENARIOS_DAMAGE" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Участок",
        "Наиболее опасный",
        "С2",
        "100,0",
        "10,0",
        "20,0",
        "30,0",
        "40,0",
        "200,0",
    ]
    properties = document.tables[0].rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None


def test_conclusions_compare_dangerous_and_probable_scenarios(
    tmp_path: Path,
) -> None:
    write_json(
        tmp_path / "risk_results.json",
        {
            "results": [
                row("С1", "Участок", 1, 3, 1200.0, 4e-5),
                row("С2", "Участок", 2, 4, 900.0, 2e-5),
                row("С3", "Участок", 0, 1, 500.0, 6e-5),
            ]
        },
    )

    assert load_key_scenario_conclusions(tmp_path) == (
        "Для составляющей ОПО «Участок» наиболее опасным является сценарий С2: "
        "погибло 2 чел., ранено 4 чел., всего пострадало 6 чел., "
        "суммарный ущерб — 900,0 тыс. руб. "
        "Наиболее вероятным является сценарий С3 с частотой 6.000E-05 1/год.",
    )


def test_conclusion_replaces_marker_with_formatted_paragraph() -> None:
    document = Document()
    document.add_paragraph("{{TOP_SCENARIOS_FINAL_CONCLUSION}}")
    conclusions = (
        "Для составляющей ОПО «Участок» наиболее опасным и наиболее вероятным "
        "является сценарий С1.",
    )

    assert render_key_scenario_conclusions(document, conclusions)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "TOP_SCENARIOS_FINAL_CONCLUSION" not in text
    assert conclusions[0] in text
    paragraph = next(
        paragraph for paragraph in document.paragraphs if paragraph.text == conclusions[0]
    )
    assert paragraph.alignment == 3
    assert paragraph.runs[0].font.name == "Times New Roman"
    assert paragraph.runs[0].font.size.pt == 10


def test_accident_description_table_has_three_people_counts() -> None:
    document = Document()
    document.add_paragraph("{{SITUATION_PLAN_ACCIDENTS_TABLE}}")
    rows = (
        {
            "component": "Участок",
            "scenario_type": "Наиболее опасный",
            "scenario_code": "С2",
            "equipment": "Нефтепровод",
            "description": "Разрыв трубопровода → пожар пролива",
            "frequency": "2.000E-05",
            "accident_mass": "4,321",
            "zones": "зона теплового излучения — 40,0 м",
            "method": "Приказ МЧС России от 10.07.2009 № 404",
            "people": "Пострадавшие: 6\nРаненые: 4\nПогибшие: 2",
            "damage": "900,0",
        },
    )

    assert render_accident_description_table(document, rows)
    assert len(document.tables[0].columns) == 11
    assert document.tables[0].rows[1].cells[9].text == (
        "Пострадавшие: 6\nРаненые: 4\nПогибшие: 2"
    )


def test_accident_description_rows_combine_current_calculations(
    tmp_path: Path,
) -> None:
    risk = row("С1", "Участок", 2, 4, 900.0, 2e-5)
    write_json(tmp_path / "risk_results.json", {"results": [risk]})
    common = {
        "id": 1,
        "scenario_code": "С1",
        "equipment_name": "Оборудование С1",
        "hazard_component": "Участок",
        "calc_code": 1,
        "ov_in_accident_t": 4.321,
    }
    factor = dict(common, ov_in_hazard_factor_t=1.234)
    write_json(tmp_path / "release_results.json", {"results": [common]})
    write_json(tmp_path / "hazard_factor_results.json", {"results": [factor]})
    impact = dict(common, impact_values={"q_10_5_m": 10.0})
    write_json(tmp_path / "impact_zones.json", {"results": [impact]})

    rows = load_accident_description_rows(tmp_path)

    assert len(rows) == 2
    assert rows[0]["scenario_code"] == "С1"
    assert rows[0]["accident_mass"] == "4,321"
    assert rows[0]["method"] == "Приказ МЧС России от 10.07.2009 № 404"
    assert rows[0]["people"] == "Пострадавшие: 6\nРаненые: 4\nПогибшие: 2"
