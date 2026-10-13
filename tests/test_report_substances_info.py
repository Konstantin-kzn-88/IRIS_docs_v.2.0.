import json
from pathlib import Path

from docx import Document

from iris_v2.report_substances_info import (
    render_substances_info_section,
    short_substance_characteristic,
)


def substance() -> dict:
    return {
        "name": "Нефть",
        "kind": 0,
        "notes": "Горючая жидкость",
        "physical": {"density_liquid_kg_per_m3": 850},
        "explosion": {"flash_point_C": -35},
        "toxicity": {"hazard_class": 3, "pdk_mg_per_m3": 100},
        "impact": "Пары раздражают органы дыхания",
        "neutralization_methods": "Собрать механическим способом",
        "first_aid": "Вывести пострадавшего на свежий воздух",
    }


def test_short_characteristic_contains_properties_and_human_impact() -> None:
    text = short_substance_characteristic(substance())

    assert "Легковоспламеняющаяся жидкость" in text
    assert "плотность жидкости 850 кг/м³" in text
    assert "температура вспышки -35 °C" in text
    assert "Воздействие на людей: Пары раздражают органы дыхания" in text
    assert "Способы обезвреживания: Собрать механическим способом" in text
    assert "Первая помощь: Вывести пострадавшего на свежий воздух" in text


def test_associated_gas_uses_gas_density_range() -> None:
    path = (
        Path(__file__).parents[1]
        / "substances"
        / "archive"
        / "DNS_UPSV"
        / "41_neft_gas.json"
    )
    item = json.loads(path.read_text(encoding="utf-8"))

    text = short_substance_characteristic(item)

    assert text.startswith(
        "Горючий газ. Основные свойства: плотность газа 1,1-4,5 кг/м³"
    )
    assert "плотность жидкости 450 кг/м³" not in text


def test_table_replaces_marker_and_has_two_columns() -> None:
    document = Document()
    document.add_paragraph("{{SUBSTANCES_INFO_SECTION}}")

    assert render_substances_info_section(document, (substance(),))
    assert "SUBSTANCES_INFO_SECTION" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    table = document.tables[0]
    assert len(table.columns) == 2
    assert [cell.text for cell in table.rows[0].cells] == [
        "Наименование вещества",
        "Краткая характеристика",
    ]
    assert table.rows[1].cells[0].text == "Нефть"
