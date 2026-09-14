import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

from iris_v2.report_substances_by_component import (
    ReportSubstancesByComponentError,
    load_substances_by_component_rows,
    render_substances_by_component_table,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def equipment(
    equipment_id: int,
    substance_id: int,
    component: str,
    equipment_type: int,
    count: int | None,
) -> dict:
    return {
        "id": equipment_id,
        "substance_id": substance_id,
        "hazard_component": component,
        "equipment_type": equipment_type,
        "equipment_count": count,
    }


def test_substances_are_aggregated_by_identification_characteristics(
    tmp_path: Path,
) -> None:
    write_json(
        tmp_path / "substances.json",
        [
            {"id": 1, "name": "Нефть", "kind": 9},
            {
                "id": 2,
                "name": "Попутный нефтяной газ",
                "kind": 2,
                "identification": {"environmental": True},
            },
        ],
    )
    write_json(
        tmp_path / "equipments.json",
        [
            equipment(1, 1, "Участок подготовки", 9, None),
            equipment(2, 1, "Склад реагентов", 1, 2),
            equipment(3, 2, "Сепарационная площадка", 1, 3),
        ],
    )
    write_json(
        tmp_path / "amount_results.json",
        {
            "results": [
                {"equipment_id": 1, "amount_t": 5.0},
                {"equipment_id": 2, "amount_t": 10.0},
                {"equipment_id": 3, "amount_t": 1.25},
            ]
        },
    )

    assert load_substances_by_component_rows(tmp_path) == (
        {
            "substance": "Нефть",
            "total": "25",
            "individual": "0",
            "flammable_gas": "0",
            "flammable_liquid_storage": "20",
            "flammable_liquid_process": "5",
            "toxic": "0",
            "highly_toxic": "0",
            "oxidizing": "0",
            "explosive": "0",
            "environmental": "0",
        },
        {
            "substance": "Попутный нефтяной газ",
            "total": "3,75",
            "individual": "0",
            "flammable_gas": "3,75",
            "flammable_liquid_storage": "0",
            "flammable_liquid_process": "0",
            "toxic": "0",
            "highly_toxic": "0",
            "oxidizing": "0",
            "explosive": "0",
            "environmental": "3,75",
        },
    )


def test_extra_amount_result_is_rejected(tmp_path: Path) -> None:
    write_json(tmp_path / "substances.json", [{"id": 1, "name": "Нефть"}])
    write_json(
        tmp_path / "equipments.json",
        [equipment(1, 1, "Участок", 9, None)],
    )
    write_json(
        tmp_path / "amount_results.json",
        {
            "results": [
                {"equipment_id": 1, "amount_t": 5.0},
                {"equipment_id": 2, "amount_t": 1.0},
            ]
        },
    )

    with pytest.raises(ReportSubstancesByComponentError, match="отсутствующее"):
        load_substances_by_component_rows(tmp_path)


def test_table_replaces_marker_and_repeats_header() -> None:
    document = Document()
    document.add_paragraph("Таблица 1 – Данные о количестве опасных веществ")
    document.add_paragraph("{{SUBSTANCES_BY_COMPONENT_TABLE}}")
    rows = (
        {
            "substance": "Нефть",
            "total": "25",
            "individual": "0",
            "flammable_gas": "0",
            "flammable_liquid_storage": "0",
            "flammable_liquid_process": "25",
            "toxic": "0",
            "highly_toxic": "0",
            "oxidizing": "0",
            "explosive": "0",
            "environmental": "0",
        },
    )

    assert render_substances_by_component_table(document, rows)
    assert "SUBSTANCES_BY_COMPONENT" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    table = document.tables[0]
    assert table.cell(0, 0).text == "Вещество"
    assert table.cell(0, 2).text == "Признаки идентификации"
    assert [cell.text for cell in table.rows[3].cells] == [
        "Нефть",
        "25",
        "–",
        "–",
        "–",
        "25",
        "–",
        "–",
        "–",
        "–",
        "–",
    ]
    assert [cell.text for cell in table.rows[4].cells] == [
        "Всего на ОПО:", "25", "–", "–", "–", "25", "–", "–", "–", "–", "–"
    ]
    properties = table.rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None
    assert table.cell(1, 2)._tc.tcPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}textDirection"
    ) is not None
    assert table.rows[3].cells[0].paragraphs[0].runs[0].font.size.pt == 11
    assert [section.orientation for section in document.sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
        WD_ORIENT.PORTRAIT,
    ]
