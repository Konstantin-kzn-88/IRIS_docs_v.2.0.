import json
from pathlib import Path

from docx import Document

from iris_v2.report_comparative_fatality_risk import (
    SOURCE_NOTE,
    load_comparative_fatality_risk_rows,
    render_comparative_fatality_risk_table,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def prepare_project(tmp_path: Path, individual: float | None) -> None:
    write_json(
        tmp_path / "risk_results.json",
        {
            "results": [
                {
                    "scenario_code": "С1",
                    "hazard_component": "Трубопроводы",
                    "fatalities_count": 1,
                    "scenario_frequency": 6e-5,
                    "collective_risk_fatalities": 6e-5,
                    "individual_risk_fatalities": individual,
                }
            ]
        },
    )
    write_json(
        tmp_path / "risk_summary.json",
        {
            "case_count": 1,
            "component_count": 1,
            "risk_unit": "1/год",
            "components": [
                {
                    "hazard_component": "Трубопроводы",
                    "scenario_count": 1,
                    "collective_risk_fatalities": 6e-5,
                    "individual_risk_fatalities": individual,
                    "fatal_accident_frequency": 6e-5,
                }
            ],
        },
    )


def test_rows_compare_component_with_background_risk(tmp_path: Path) -> None:
    prepare_project(tmp_path, 4e-6)

    assert load_comparative_fatality_risk_rows(tmp_path) == (
        {
            "component": "Трубопроводы",
            "risk": "4.000E-06",
            "ppm": "4,000",
            "dbr": "-16,9",
            "comparison": "Ниже в 48,8 раза",
        },
        {
            "component": "Фоновый риск RГЛ",
            "risk": "1.950E-04",
            "ppm": "195",
            "dbr": "0,0",
            "comparison": "Фоновое значение",
        },
    )


def test_missing_individual_risk_is_not_invented(tmp_path: Path) -> None:
    prepare_project(tmp_path, None)

    row = load_comparative_fatality_risk_rows(tmp_path)[0]

    assert row["risk"] == "Не рассчитан"
    assert row["ppm"] == "—"
    assert row["dbr"] == "—"
    assert row["comparison"] == "Не рассчитан"


def test_zero_risk_uses_negative_infinity(tmp_path: Path) -> None:
    prepare_project(tmp_path, 0.0)

    row = load_comparative_fatality_risk_rows(tmp_path)[0]

    assert row["ppm"] == "0,000"
    assert row["dbr"] == "−∞"
    assert row["comparison"] == "Ниже фонового значения"


def test_table_replaces_marker_and_adds_source_note() -> None:
    document = Document()
    document.add_paragraph("{{COMPARATIVE_FATALITY_RISK_TABLE}}")
    rows = (
        {
            "component": "Трубопроводы",
            "risk": "4.000E-06",
            "ppm": "4,000",
            "dbr": "-16,9",
            "comparison": "Ниже в 48,8 раза",
        },
        {
            "component": "Фоновый риск RГЛ",
            "risk": "1.950E-04",
            "ppm": "195",
            "dbr": "0,0",
            "comparison": "Фоновое значение",
        },
    )

    assert render_comparative_fatality_risk_table(document, rows)
    assert "COMPARATIVE_FATALITY" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert SOURCE_NOTE in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Трубопроводы",
        "4.000E-06",
        "4,000",
        "-16,9",
        "Ниже в 48,8 раза",
    ]
