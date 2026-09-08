import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from iris_v2.report_component_inputs_assumptions import (
    load_component_inputs_assumptions,
    render_component_inputs_assumptions_section,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def test_component_inputs_and_assumptions_are_collected(tmp_path: Path) -> None:
    write_json(
        tmp_path / "substances.json",
        [{"id": 1, "name": "Нефть"}],
    )
    write_json(
        tmp_path / "equipments.json",
        [
            {
                "id": 1,
                "substance_id": 1,
                "equipment_name": "Нефтепровод № 1",
                "equipment_type": 0,
                "phase_state": "ж.ф.",
                "pressure_mpa": 1.6,
                "substance_temperature_c": 20.0,
                "shutdown_time_s": 12.0,
                "evaporation_time_s": 3600.0,
                "spill_coefficient": 20.0,
                "spill_area_m2": 0.0,
                "clutter_degree": 2,
                "hazard_component": "Участок трубопроводов",
            }
        ],
    )
    write_json(
        tmp_path / "amount_results.json",
        {"results": [{"equipment_id": 1, "amount_t": 5.25}]},
    )
    write_json(
        tmp_path / "calculation_cases.json",
        {
            "cases": [
                {
                    "scenario_code": "С1",
                    "equipment_id": 1,
                    "scenario_text": "Разрыв → пожар пролива",
                }
            ]
        },
    )

    rows, assumptions = load_component_inputs_assumptions(tmp_path)

    assert rows == (
        {
            "component": "Участок трубопроводов",
            "equipment": "Нефтепровод № 1",
            "substance_mass": "Нефть; 5,25 т",
            "regime": "ж.ф.; P=1,6 МПа; T=20 °C",
            "scenarios": "С1: Разрыв → пожар пролива",
            "assumptions": (
                "отключение: 12 с; испарение: 3600 с; растекание: 20 м⁻¹; "
                "площадь пролива: 0 м²; загромождённость: 2"
            ),
        },
    )
    assert ("Скорость ветра", "1 м/с") in assumptions
    assert (
        "Множители частот: стандартный / без КМ / с КМ",
        "1 / 1,25 / 0,6",
    ) in assumptions


def test_component_inputs_tables_have_no_fill() -> None:
    document = Document()
    document.add_paragraph("{{COMPONENT_INPUTS_ASSUMPTIONS_SECTION}}")
    rows = (
        {
            "component": "Участок",
            "equipment": "Трубопровод",
            "substance_mass": "Нефть; 5 т",
            "regime": "ж.ф.; P=1 МПа; T=20 °C",
            "scenarios": "С1: Разрыв",
            "assumptions": "отключение: 12 с",
        },
    )

    assert render_component_inputs_assumptions_section(
        document,
        rows,
        (("Скорость ветра", "1 м/с"),),
    )
    assert len(document.tables) == 2
    assert all(
        table._tbl.findall(f".//{qn('w:shd')}") == []
        for table in document.tables
    )
