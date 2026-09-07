import json
from pathlib import Path

from docx import Document

from iris_v2.report_ngk_background_risk import (
    SOURCE_NOTE,
    load_ngk_background_risk_rows,
    render_ngk_background_risk_comparison,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def prepare_project(tmp_path: Path, individual: float | None) -> None:
    write_json(
        tmp_path / "risk_results.json",
        {
            "results": [
                {
                    "scenario_code": "С1",
                    "hazard_component": "Трубопроводы",
                    "fatalities_count": 1,
                    "scenario_frequency": 6e-5,
                    "collective_risk_fatalities": 6e-5,
                    "individual_risk_fatalities": individual,
                }
            ]
        },
    )
    write_json(
        tmp_path / "risk_summary.json",
        {
            "case_count": 1,
            "component_count": 1,
            "risk_unit": "1/год",
            "components": [
                {
                    "hazard_component": "Трубопроводы",
                    "scenario_count": 1,
                    "collective_risk_fatalities": 6e-5,
                    "individual_risk_fatalities": individual,
                    "fatal_accident_frequency": 6e-5,
                }
            ],
        },
    )


def test_rows_compare_maximum_risk_with_all_ngk_industries(tmp_path: Path) -> None:
    prepare_project(tmp_path, 4e-6)

    maximum, rows = load_ngk_background_risk_rows(tmp_path)

    assert maximum == "4.000E-06"
    assert rows == (
        {
            "industry": "Нефтегазодобывающая промышленность",
            "dbr": "-4,3",
            "per_100k": "7,3",
            "per_year": "7.300E-05",
            "comparison": "Ниже в 18,2 раза",
        },
        {
            "industry": "Нефтеперерабатывающая промышленность",
            "dbr": "-5,8",
            "per_100k": "5,2",
            "per_year": "5.200E-05",
            "comparison": "Ниже в 13,0 раза",
        },
        {
            "industry": "Нефтехимическая промышленность",
            "dbr": "-8,9",
            "per_100k": "2,4",
            "per_year": "2.400E-05",
            "comparison": "Ниже в 6,0 раза",
        },
        {
            "industry": "Объекты газораспределения и газопотребления*",
            "dbr": "-8,7",
            "per_100k": "2,6",
            "per_year": "2.600E-05",
            "comparison": "Ниже в 6,5 раза",
        },
        {
            "industry": "Магистральный трубопроводный транспорт*",
            "dbr": "-10,9",
            "per_100k": "1,6",
            "per_year": "1.600E-05",
            "comparison": "Ниже в 4,0 раза",
        },
    )


def test_missing_individual_risk_is_not_compared(tmp_path: Path) -> None:
    prepare_project(tmp_path, None)

    maximum, rows = load_ngk_background_risk_rows(tmp_path)

    assert maximum == "Не рассчитан"
    assert all(row["comparison"] == "Не рассчитан" for row in rows)


def test_section_replaces_marker_and_adds_source_note() -> None:
    document = Document()
    document.add_paragraph("{{NGK_BACKGROUND_RISK_COMPARISON}}")
    rows = (
        {
            "industry": "Нефтегазодобывающая промышленность",
            "dbr": "-4,3",
            "per_100k": "7,3",
            "per_year": "7.300E-05",
            "comparison": "Ниже в 18,2 раза",
        },
    )

    assert render_ngk_background_risk_comparison(
        document, "4.000E-06", rows
    )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "NGK_BACKGROUND" not in text
    assert "4.000E-06 1/год" in text
    assert SOURCE_NOTE in text
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Нефтегазодобывающая промышленность",
        "-4,3",
        "7,3",
        "7.300E-05",
        "Ниже в 18,2 раза",
    ]
