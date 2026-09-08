import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from iris_v2.report_component_risk_summary import (
    load_component_risk_summary_rows,
    render_component_risk_summary_table,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def test_component_risk_summary_contains_all_people_categories(
    tmp_path: Path,
) -> None:
    results = [
        {
            "scenario_code": "С1",
            "hazard_component": "Участок А",
            "scenario_frequency": 2e-5,
            "fatalities_count": 2,
            "injured_count": 3,
            "total_damage": 1500.0,
            "collective_risk_fatalities": 4e-5,
            "individual_risk_fatalities": 2e-6,
        },
        {
            "scenario_code": "С2",
            "hazard_component": "Участок А",
            "scenario_frequency": 3e-5,
            "fatalities_count": 1,
            "injured_count": 7,
            "total_damage": 2500.0,
            "collective_risk_fatalities": 3e-5,
            "individual_risk_fatalities": 1.5e-6,
        },
    ]
    write_json(tmp_path / "risk_results.json", {"results": results})
    write_json(
        tmp_path / "risk_summary.json",
        {
            "case_count": 2,
            "component_count": 1,
            "risk_unit": "1/год",
            "damage_unit": "тыс. руб.",
            "components": [
                {
                    "hazard_component": "Участок А",
                    "scenario_count": 2,
                    "max_total_damage": 2500.0,
                    "collective_risk_fatalities": 7e-5,
                    "individual_risk_fatalities": 3.5e-6,
                }
            ],
        },
    )

    assert load_component_risk_summary_rows(tmp_path) == (
        {
            "component": "Участок А",
            "frequency": "5.000E-05",
            "fatalities": "2",
            "injured": "7",
            "affected": "8",
            "damage": "2500,0",
            "collective": "7.000E-05",
            "individual": "3.500E-06",
        },
    )


def test_component_risk_table_has_no_fill(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("{{COMPONENT_RISK_SUMMARY_TABLE}}")
    rows = (
        {
            "component": "Участок А",
            "frequency": "1.000E-05",
            "fatalities": "1",
            "injured": "2",
            "affected": "3",
            "damage": "100,0",
            "collective": "1.000E-05",
            "individual": "1.000E-06",
        },
    )

    assert render_component_risk_summary_table(document, rows)
    assert document.tables[0]._tbl.findall(f".//{qn('w:shd')}") == []
