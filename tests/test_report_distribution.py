from docx import Document
from docx.enum.section import WD_ORIENT

from iris_v2.report_distribution import render_distribution_section


def test_dpb_distribution_table_uses_own_landscape_section() -> None:
    document = Document()
    document.add_paragraph(
        "Таблица 5 – Данные о распределении опасных веществ по оборудованию"
    )
    document.add_paragraph("{{DISTRIBUTION_SECTION}}")
    document.add_paragraph("Следующий раздел")
    equipment = (
        {
            "id": 1,
            "equipment_name": "Нефтепровод",
            "hazard_component": "Участок трубопроводов",
            "substance_id": 1,
            "equipment_type": 0,
            "pressure_mpa": 1.6,
            "substance_temperature_c": 20.0,
            "phase_state": "ж.ф.",
        },
    )
    substances = ({"id": 1, "name": "Нефть"},)
    amounts = ({"equipment_id": 1, "amount_t": 2.5},)

    assert render_distribution_section(document, equipment, substances, amounts)
    assert [section.orientation for section in document.sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
        WD_ORIENT.PORTRAIT,
    ]
    assert document.paragraphs[-1].text == "Следующий раздел"
