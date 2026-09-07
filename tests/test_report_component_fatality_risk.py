import json
from pathlib import Path

import pytest
from docx import Document

from iris_v2.report_component_fatality_risk import (
    ReportComponentFatalityRiskError,
    load_component_fatality_risk_rows,
    render_component_fatality_risk_section,
)


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def risk_row(
    code: str,
    component: str,
    fatalities: int,
    frequency: float,
    collective: float,
    individual: float | None,
) -> dict:
    return {
        "scenario_code": code,
        "hazard_component": component,
        "fatalities_count": fatalities,
        "scenario_frequency": frequency,
        "collective_risk_fatalities": collective,
        "individual_risk_fatalities": individual,
    }


def component(
    name: str,
    count: int,
    collective: float,
    individual: float | None,
    fatal_frequency: float,
) -> dict:
    return {
        "hazard_component": name,
        "scenario_count": count,
        "collective_risk_fatalities": collective,
        "individual_risk_fatalities": individual,
        "fatal_accident_frequency": fatal_frequency,
    }


def test_fatality_risk_is_grouped_by_component(tmp_path: Path) -> None:
    results = [
        risk_row("С1", "Трубопроводы", 2, 2e-5, 4e-5, 2e-6),
        risk_row("С2", "Резервуарный парк", 0, 7e-5, 0.0, 0.0),
        risk_row("С3", "Трубопроводы", 1, 3e-5, 3e-5, 1.5e-6),
    ]
    write_json(tmp_path / "risk_results.json", {"results": results})
    write_json(
        tmp_path / "risk_summary.json",
        {
            "case_count": 3,
            "component_count": 2,
            "risk_unit": "1/год",
            "components": [
                component("Трубопроводы", 2, 7e-5, 3.5e-6, 5e-5),
                component("Резервуарный парк", 1, 0.0, 0.0, 0.0),
            ],
        },
    )

    assert load_component_fatality_risk_rows(tmp_path) == (
        {
            "component": "Трубопроводы",
            "max_fatalities": "2",
            "collective": "7.000E-05",
            "individual": "3.500E-06",
            "fatal_frequency": "5.000E-05",
        },
        {
            "component": "Резервуарный парк",
            "max_fatalities": "0",
            "collective": "0.000E+00",
            "individual": "0.000E+00",
            "fatal_frequency": "0.000E+00",
        },
    )


def test_missing_people_is_shown_as_not_calculated(tmp_path: Path) -> None:
    results = [risk_row("С1", "Трубопроводы", 1, 1e-5, 1e-5, None)]
    write_json(tmp_path / "risk_results.json", {"results": results})
    write_json(
        tmp_path / "risk_summary.json",
        {
            "case_count": 1,
            "component_count": 1,
            "risk_unit": "1/год",
            "components": [component("Трубопроводы", 1, 1e-5, None, 1e-5)],
        },
    )

    assert load_component_fatality_risk_rows(tmp_path)[0]["individual"] == (
        "Не рассчитан"
    )


def test_stale_fatality_risk_is_rejected(tmp_path: Path) -> None:
    results = [risk_row("С1", "Трубопроводы", 1, 1e-5, 1e-5, 1e-6)]
    write_json(tmp_path / "risk_results.json", {"results": results})
    write_json(
        tmp_path / "risk_summary.json",
        {
            "case_count": 1,
            "component_count": 1,
            "risk_unit": "1/год",
            "components": [component("Трубопроводы", 1, 1e-5, 1e-6, 2e-5)],
        },
    )

    with pytest.raises(ReportComponentFatalityRiskError, match="устарел"):
        load_component_fatality_risk_rows(tmp_path)


def test_section_replaces_marker_with_repeatable_table() -> None:
    document = Document()
    document.add_paragraph("{{FATALITY_RISK_BY_COMPONENT_SECTION}}")
    rows = (
        {
            "component": "Трубопроводы",
            "max_fatalities": "2",
            "collective": "7.000E-05",
            "individual": "3.500E-06",
            "fatal_frequency": "5.000E-05",
        },
    )

    assert render_component_fatality_risk_section(document, rows)
    assert "FATALITY_RISK" not in "\n".join(p.text for p in document.paragraphs)
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "Трубопроводы",
        "2",
        "7.000E-05",
        "3.500E-06",
        "5.000E-05",
    ]
    properties = document.tables[0].rows[0]._tr.get_or_add_trPr()
    assert properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
    ) is not None
